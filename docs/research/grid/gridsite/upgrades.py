"""Phase 4: which transmission upgrades would change where batteries should go, what they cost, and could a battery do the job.

Four outputs (all CSVs have a `.labels.csv` sidecar; labels are REAL / DERIVED / SIM / ASSUMPTION as in PLAN.md):

  data/upgrades.csv                  one row per monitored element: type, kV, length, the primary wires upgrade with a low / mid /
                                     high cost and a lead time, planned-project match, non-wires (battery) sizing summary, who decides
  data/upgrade_options.csv           every plausible wires option per element (reconductor, rebuild, new parallel line, dynamic /
                                     ambient-adjusted ratings, autotransformer, 765 kV plan) with the cost formula
  data/nwa_sizing.csv                battery MW / MWh that would relieve each element: 3 relief targets x 3 coverage targets x 2 sides
  data/upgrade_rerank_first_order.csv  per upgrade, the sites that lose congestion value and the sites that rank up (first order, no dispatch)

Method in one paragraph. The market data cannot say how much relief is "needed" (SCED keeps flow at or below the limit), so relief
is a named scenario (5 / 10 / 20 % of the median limit, ASSUMPTION). Coverage (50 / 80 / 95 % of binding hours) sets the energy: binding
episodes are read from the RT and DAM windows (union over the element's contingencies; gaps up to EPISODE_GAP_MIN merge), and the
duration D(x) is the smallest D such that sustaining relief for at most D hours from the start of each episode covers x of all binding
hours. MW = relief / |SF| at the best site on the side that relieves (import side by discharging, export side by charging), site by
site, with a per-site MW cap. Battery cost is the NREL 2025 update (2024 dollars): energy and power cost lines, 2 h and 4 h and any
duration between. Wires cost is the MISO MTEP24 Texas unit-cost rows in data/unit_costs.csv times a quantity (miles, MVA, bays), with
the formula in the table. The first-order rerank values a site at |SF[n,k] - zone mean SF[k]| x annual shadow-price-hours[k] x
CAPTURE_FACTOR and sums over constraints; an upgrade scales the shadow prices of its constraints by a residual (apply_upgrade).

Run (from docs/research/grid, research venv):
  run_grid.py upgrades              data/upgrades.csv + data/upgrade_options.csv
  run_grid.py nwa                   data/nwa_sizing.csv
  run_grid.py rerank-first-order    data/upgrade_rerank_first_order.csv
  run_grid.py upgrades check        recompute the derived constants (capture cover, lead times, SF fallback) and compare with the named ones

Interface for Phase 3b and the final report:
  apply_upgrade(cid_list, residual=None, expand_element=True) -> {cid: mu_scale}      0.0 = fully relieved
  planned_schedule()                -> DataFrame of (cid, project, in-service date, status, confidence) for the planned-upgrade clock
  basis_half_life(v_by_cid, ...)    -> years until a site's congestion value halves given the planned schedule
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from functools import lru_cache

import numpy as np
import pandas as pd

from . import paths

# =====================================================================================================================
# Constants. Every unverified value is a single named constant with its label (REAL / DERIVED / SIM / ASSUMPTION) and source.
# =====================================================================================================================

# --- dates, units ---------------------------------------------------------------------------------------------------
AS_OF = pd.Timestamp("2026-10-10")        # REAL: date of this build (system date); planned dates before it are "imminent"
USD_YEAR = 2024                           # REAL: price year of the MISO MTEP24 rows (unit_costs.csv) and of the NREL 2025 battery update
ESCALATION_PER_YEAR = 0.04                # ASSUMPTION: MISO escalates ~4%/yr (unit_costs.labels.csv); NOT applied, all dollars stay 2024$
KM_PER_MILE = 1.609344                    # REAL: unit conversion
EARTH_RADIUS_KM = 6371.0088               # REAL: mean Earth radius
ROUTE_FACTOR = 1.25                       # ASSUMPTION: routed length / straight-line length (same value as gridsite.costs.ROUTE_FACTOR)
SCED_DT_MIN = 5.0                         # ASSUMPTION: one SCED run covers 5 minutes (same convention as phase 1a / phase 2)
DAM_DT_MIN = 60.0                         # REAL: DAM prices are hourly
UPGRADE_CONF_MIN = 0.5                    # ASSUMPTION: a planned-upgrade match at or above this counts as "has a planned upgrade" (phase 2)
PLANNED_YES_CONF = 0.8                    # ASSUMPTION: match confidence at or above this is called "yes" (pair_name 0.9, pair_title 0.8)
PLANNED_HINT_CONF = 0.3                   # ASSUMPTION: a match at or above this (phase 2 one_end, 0.30) is a "hint": a nearby project, not the element
ANCHOR_CONF = 0.5                         # ASSUMPTION: substation placement confidence needed for a station-distance length (phase 2 / 1c)

# --- element selection ----------------------------------------------------------------------------------------------
TOP_PRIORITY_ELEMENTS = 15                # ASSUMPTION: brief says "top ~15 binding elements by constraints.csv priority"
WINDOW_TOP_ELEMENTS = 6                   # ASSUMPTION: extra local constraints, by the window part of the priority score, not yet chosen
# Every element with a 2025 IMM top-10 rent (constraints.csv annual_rent_2025_musd) is included even if it never bound in the window.

# --- lengths, limits, binding hours (fallbacks, all labelled in the output) ------------------------------------------
LENGTH_FALLBACK_QUANTILE = 0.5            # ASSUMPTION: class median line length (lines.csv, OSM, both ends ERCOT substations) when nothing better
LIMIT_FALLBACK_QUANTILE = 0.5             # ASSUMPTION: class median of median_limit_mw over window-binding elements of the same type and kV class
WINDOW_PERSISTENCE = 0.15                 # ASSUMPTION: share of the window's binding that recurs over a year (outage-driven shoulder week);
                                          #   Scurry Chevron-Knapp, the one element with both, measures 0.17 (`upgrades check`)
IMM_FLOOR_RANK_RENT_MUSD = 50.83          # REAL: 10th-largest 2025 IMM RT rent (Argyle-Lewisville, constraints_annual.csv); an element outside the
                                          #   IMM top 10 has less annual rent than this, so it caps window-annualised rents

# --- episodes (binding-hour structure) ------------------------------------------------------------------------------
EPISODE_MIN_SHADOW = 1.0                  # ASSUMPTION: $/MWh; a shadow price below this is not material (<0.5% of rent in the window; 33% of DAM rows)
EPISODE_GAP_MIN = 60.0                    # ASSUMPTION: binding periods separated by at most this many minutes belong to one episode (no time to recharge)
EPISODE_TOL_MIN = 1.5                     # ASSUMPTION: SCED runs are 295-305 s apart; this much slack when testing run adjacency
MIN_BASIS_EPISODES = 6                    # ASSUMPTION: an element needs this many window episodes to use its own episode lengths
MIN_POOL_EPISODES = 20                    # ASSUMPTION: a pooled reference class needs this many episodes, else all lines are pooled
COVERAGE_TARGETS = (0.50, 0.80, 0.95)     # REAL: from the brief (50 / 80 / 95 % of binding hours)
RELIEF_TARGETS = (0.05, 0.10, 0.20)       # ASSUMPTION: from the brief (named scenarios: 5 / 10 / 20 % of the median limit)
CENTRAL_RELIEF = 0.10                     # ASSUMPTION: the central NWA scenario quoted in upgrades.csv
CENTRAL_COVERAGE = 0.80                   # ASSUMPTION: the central NWA scenario quoted in upgrades.csv

# --- shift factors and battery siting for the NWA --------------------------------------------------------------------
SF_NWA_CAP = 1.0                          # ASSUMPTION: |SF| above 1 is not credited (a battery cannot relieve more than its own MW)
SF_SITE_MIN = 0.05                        # ASSUMPTION: sites with a smaller |SF| are not used to relieve (MW would explode)
SHRINK_MIN = 0.8                          # REAL: phase 2 SHRINK_MIN; DAM SFs below this are attenuated ("+shrunk")
UNSHRINK_FLOOR = 0.3                      # ASSUMPTION: an attenuated DAM SF is divided by its shrink factor if shrink >= this, else sign-only
SITE_MAX_MW = {69: 50.0, 138: 300.0, 345: 300.0}   # ASSUMPTION: MW cap per battery site by voltage class (brief: 10-300 MW at 69-345 kV; 69 kV
                                                   #   lines are rated ~40 MW at the median in the window, so 50 MW there)
BAD_SF_QUALITY = ("weak", "implausible", "low")    # ASSUMPTION: a site SF with any of these quality tags is not used to size relief
FALLBACK_SF_QUANTILE = 0.5                # ASSUMPTION: no regression SF -> use this quantile of the best-site |SF| of window-fitted RT lines of the same kV class

# --- battery cost (NREL, 2024 dollars) -----------------------------------------------------------------------------
# Cole, Ramasamy, Turan (2025), "Cost Projections for Utility-Scale Battery Storage: 2025 Update", NREL/TP-6A40-93281, June 2025 (the cost
# source of the NREL ATB 2025 utility-scale battery page). Fig. 2: total ($/kW) = energy cost ($/kWh) x hours + power cost ($/kW), 60 MW
# system, Q1 2024 modelled market price. 4 h = $334/kWh = $1,336/kW.
BATT_ENERGY_USD_KWH = 241.0               # REAL: NREL/TP-6A40-93281 Fig. 2 (fit slope 240.8), 2024$
BATT_POWER_USD_KW = 372.0                 # REAL: NREL/TP-6A40-93281 Fig. 2 (fit intercept 379.2 less the 7 $/kW adjustment to hit $334/kWh at 4 h), 2024$
BATT_COST_MULT = {"low": 0.77, "mid": 0.92, "high": 1.10}   # REAL: NREL/TP-6A40-93281 Table 2, normalised 2026 column (low / mid / high), applied to both lines
BATT_FOM_FRAC = 0.04                      # REAL: NREL/TP-6A40-93281 sec. 3: fixed O&M = 4 %/yr of capital cost (includes augmentation)
BATT_LIFE_YR = 15                         # REAL: NREL/TP-6A40-93281 sec. 3 (15 years)
BATT_RTE = 0.85                           # REAL: NREL/TP-6A40-93281 sec. 3 (round trip 85 %); informational, usable MWh is quoted
BATT_AVAILABILITY = 0.95                  # ASSUMPTION: installed MW = required MW / availability (forced outages; relief is only worth what is available)
BATT_MIN_HOURS = 1.0                      # ASSUMPTION: floor on duration (NREL cautions against < 1 h; ERCOT ECRS needs 1 h)
BATT_MAX_HOURS = 10.0                     # REAL: the NREL energy / power fit and the ATB cover 1-10 h; beyond it the table caps the duration and reports the coverage reached
DISCOUNT_RATE = 0.08                      # ASSUMPTION: 8 %, the rate the arena uses for break-even installed cost (arena README)
WIRES_FIXED_CHARGE_RATE = 0.115           # ASSUMPTION: levelised annual revenue requirement / installed cost (return, depreciation, taxes, O&M)
REF_ZONE_ENERGY_USD_KW_YR = 44.5          # REAL: arena README, ERCOT 5-year (2021-25) Adaptive learner, grid-scale 2 h battery; reference only, never netted
ARENA_LEARNER_SHARE = 0.41                # REAL: arena README, ERCOT learner = 41 % of the perfect-foresight ceiling (5 y; 37 % in 2025, 51 % over 3 y)

# --- wires cost model (quantities and multipliers on top of data/unit_costs.csv) ---------------------------------------
TERMINAL_BAY_FRACTION = 0.30              # ASSUMPTION: terminal equipment (breakers, CTs, switches) per line end = this x the cost of one new bay
NEW_LINE_BAYS = 2                         # ASSUMPTION: a new line needs one new bay at each end
REBUILD_345_FROM_NEW = 1.8 / 2.3          # DERIVED: no 345 kV rebuild row; 138 kV rebuild / 138 kV new single-circuit ratio (1.8 / 2.3 $M/mi) applied to 345 kV new
ESTIMATE_BAND_LOW = 0.20                  # ASSUMPTION: screening-level estimate: low >= mid x (1 - 0.20) ...
ESTIMATE_BAND_HIGH = 0.50                 # ASSUMPTION: ... and high >= mid x (1 + 0.50), widened only where the source cells are narrower
XFMR_STEPS_MVA = (150, 300, 450, 600, 750, 900, 1200)   # ASSUMPTION: standard autotransformer sizes; a parallel unit is the smallest size >= the existing rating
XFMR_138_69_USD_PER_MVA = 5050.0          # REAL (UNVERIFIED layout): MISO MTEP24 Table 2.3-7, 138 kV row of the 69 kV column; 345/138 kV is $6,880/MVA in unit_costs.csv
DLR_MIN_COST_MUSD = 0.25                  # ASSUMPTION: floor on a DLR / AAR installation (weather stations, integration), whatever the line length
SUBSTATION_REBUILD_KV_ROW = {69: 138, 138: 138, 345: 345}   # ASSUMPTION: which unit-cost row prices a substation rebuild at each voltage class
UPLIFT = {                                # ASSUMPTION: typical rating uplift of each option; an option can serve a relief target only if uplift >= target
    "dlr_aar_static": 0.15,               #   DOE 2019 DLR report Table 2: 10-20 % typical over static, 22 % average in one study
    "dlr_aar_after_aar": 0.10,            #   same table: 6-14 % over ambient-adjusted ratings (used where the SCED limit already moves with temperature)
    "reconductor": 0.40,                  #   conventional reconductor; advanced conductors can reach +100 % (DOE 2019)
    "rebuild": 0.60,
    "new_parallel_line": 1.00,
    "autotransformer": 1.00,              #   a parallel unit of the same size doubles the rating
    "substation_rebuild": 0.30,
    "plan_765kv": 0.30,                   #   the 2027 -> 2030 WESTEX planning limit rises 12,240 -> 16,200 MW (+32 %), data/gtc.csv
}
AMBIENT_SWING_MIN = 0.05                  # ASSUMPTION: if the SCED limit of an element swings >= 5 % (q95-q05)/median it is already ambient-adjusted
SWING_QUANTILES = (0.05, 0.95)            # ASSUMPTION: quantiles of the SCED limit that define the swing
SWING_MIN_ROWS = 20                       # ASSUMPTION: with fewer listed SCED rows than this the swing is not computed

# --- lead times (months from RPG submittal to in service unless stated) ---------------------------------------------
LEAD_MIN_N = 10                           # ASSUMPTION: a TPIT class needs this many rows for its own quartiles, else pooled over kV
DLR_LEAD_MONTHS = (6.0, 12.0, 18.0)       # ASSUMPTION: sensors and software, no CCN, a ratings change is not an RPG project (NP 3.11.4.1(4))
LPT_LEAD_MONTHS = (30.0, 36.0, 60.0)      # REAL: unit_costs.csv: NLR 2026 average 30 months (H1 2025), DOE 2024 commonly quoted 36, up to 60
XFMR_INSTALL_MONTHS = 6.0                 # ASSUMPTION: set, test and energise after delivery (same as gridsite.costs.XFMR_INSTALL_MONTHS)
BATT_FIS_MONTHS = (16.0, 20.5, 28.0)      # DERIVED: quartiles / median of FIS request -> IA signed for storage (notes/rules.md (d), GIS Sep 2026)
BATT_BUILD_MONTHS = (12.0, 15.0, 24.0)    # ASSUMPTION: engineering, procurement and construction after the interconnection agreement
BATT_LPT_MIN_MVA = 100.0                  # ASSUMPTION: a step-up of at least this size at 138 kV or more is a large power transformer (same as gridsite.costs.LPT_MIN_MVA)
BATT_GSU_MVA_PER_MW = 1.1                 # ASSUMPTION: step-up rating = 110 % of the battery MW (same as gridsite.costs.GSU_MVA_PER_MW)
PLAN_765_ISD = ("2030-12-31", "2031-06-30", "2032-12-31")   # REAL: PBRP 765-kV import paths "2031" (C&N Dec 2025 p.8; data/planned_projects.csv); early / mid / late

# --- first-order rerank ---------------------------------------------------------------------------------------------
DAM_WINDOW_DAYS = 31.0                    # REAL: DAM shadow prices 10 Sep - 10 Oct 2026
MIN_RT_RUNS_FOR_SPH = 10                  # REAL: phase 2 MIN_BINDING_RUNS; below this the RT window is too thin, use DAM
COVER_MIN_SPH = 200.0                     # ASSUMPTION: constraints with less shadow-price-hours than this ($/MW over the window) are left out of the cover check
COVER_2H = 0.50                           # DERIVED/ASSUMPTION: share of a constraint's shadow-price-hours a 2 h battery that cycles once a day can
                                          #   reach with perfect foresight (best 2 h of each day). Window values: RT 0.55, DAM 0.45 (value-weighted), `upgrades check`
COVER_4H = 0.76                           # DERIVED/ASSUMPTION: same for 4 h (RT 0.81, DAM 0.71)
CAPTURE_FACTOR = {"2h": ARENA_LEARNER_SHARE * COVER_2H, "4h": ARENA_LEARNER_SHARE * COVER_4H}   # ASSUMPTION: cover x learner share of the oracle (0.205 / 0.312)
V_MIN_USD_KW_YR = 0.5                     # ASSUMPTION: sites with less than this total value are not ranked (rank among noise is meaningless)
N_LOSERS = 10                             # ASSUMPTION: rows kept per upgrade and spec: biggest losers
N_RISERS = 10                             # ASSUMPTION: ... and the biggest rank gainers inside the post-upgrade top set
TOP_K = 50                                # ASSUMPTION: size of the "top set" whose membership is tracked (a site that enters or leaves the best 50 is a mover)
RISER_UNIVERSE = 200                      # ASSUMPTION: a "riser" is a site inside the post-upgrade best 200 whose rank improves; deeper ranks sit in a flat tail where a
                                          #   few dollars move a site hundreds of places
SATURATION_RELIEF = 0.20                  # ASSUMPTION: the per-kW values hold for the first MW only; sat_mw is the battery MW that already relieves 20 % of the limit

# --- upgrade effect on shadow prices (apply_upgrade) ------------------------------------------------------------------
RESIDUAL_MU_SCALE_WIRES = 0.10            # ASSUMPTION: share of shadow-price-hours left after a line / transformer upgrade (re-binding at the higher limit,
                                          #   other contingencies); 0 = fully relieved
RESIDUAL_MU_SCALE_GTC = 0.50              # ASSUMPTION: an interface keeps half its congestion after its planned relief (the arena halves the Permian premium
                                          #   after the 765 kV lines, arena/forward.py)

# --- who decides (text only; sources in notes/phase4.md) -----------------------------------------------------------------
CITE_PLAN = "PG 3.1.1.3, 3.1.2.1, 3.1.3; NP 3.11.4.1, 3.11.4.3, 3.11.4.7-9; 16 TAC 25.195(d)"
CITE_NOCCN = "16 TAC 25.101(c)(5)(B) (reconductor / rebuild on existing right of way needs no CCN); NP 3.11.4.3(1)(f)(viii) (ratings-only upgrade is a neutral Tier 4 project)"
CITE_CCN = "16 TAC 25.101(b)(3) (CCN for a new line; economic test (A)(i), reliability (A)(ii)); 80-day uncontested, 180-day critical-to-reliability, 1 year otherwise"
CITE_COST = "16 TAC 25.192 (TCOS recovered by 4CP-based postage-stamp rates; wholesale storage exempt)"

# =====================================================================================================================
# Paths
# =====================================================================================================================
CONSTRAINTS_CSV = paths.REPO_DATA / "constraints.csv"
UNIT_COSTS_CSV = paths.REPO_DATA / "unit_costs.csv"
PLANNED_CSV = paths.REPO_DATA / "planned_projects.csv"
LINES_CSV = paths.REPO_DATA / "lines.csv"
SUBSTATIONS_CSV = paths.REPO_DATA / "substations.csv"
SHIFT_FACTORS_CSV = paths.REPO_DATA / "shift_factors.csv"
SETTLEMENT_CSV = paths.REPO_DATA / "settlement_points.csv"
SF_LONG = paths.DERIVED / "sf_long.parquet"
SHADOW_SCED = paths.DERIVED / "shadow_sced.parquet"
SHADOW_DAM = paths.DERIVED / "shadow_dam.parquet"
TPIT_XLSX = paths.RAW / "docs" / "TPIT_Transmission_Project_and_Information_Tracking_2026-07.xlsx"
OUT_UPGRADES = paths.REPO_DATA / "upgrades.csv"
OUT_OPTIONS = paths.REPO_DATA / "upgrade_options.csv"
OUT_NWA = paths.REPO_DATA / "nwa_sizing.csv"
OUT_RERANK = paths.REPO_DATA / "upgrade_rerank_first_order.csv"

LABELS = ("REAL", "DERIVED", "SIM", "ASSUMPTION")


# =====================================================================================================================
# Small utilities
# =====================================================================================================================
def kv_class(kv) -> int | None:
    """69 (below 100 kV), 138 (100-200 kV), 345 (200-400 kV), 765. None if unknown."""
    try:
        v = float(kv)
    except (TypeError, ValueError):
        return None
    if not np.isfinite(v):
        return None
    if v < 100:
        return 69
    if v < 200:
        return 138
    if v < 400:
        return 345
    return 765


def _kc(v) -> int | None:
    """kV class stored in a DataFrame (float with NaN) back to int / None."""
    try:
        return None if v is None or v != v else int(v)
    except (TypeError, ValueError):
        return None


def haversine_km(lat1, lon1, lat2, lon2) -> float:
    la1, la2 = math.radians(lat1), math.radians(lat2)
    dphi, dl = la2 - la1, math.radians(lon2 - lon1)
    a = math.sin(dphi / 2) ** 2 + math.cos(la1) * math.cos(la2) * math.sin(dl / 2) ** 2
    return 2 * EARTH_RADIUS_KM * math.asin(math.sqrt(a))


def crf(rate: float, years: float) -> float:
    """Capital recovery factor."""
    if rate <= 0:
        return 1.0 / years
    g = (1 + rate) ** years
    return rate * g / (g - 1)


def fm(x, nd=1) -> str:
    """Format $M for formula strings."""
    if x is None or (isinstance(x, float) and not np.isfinite(x)):
        return "n/a"
    return f"${x:,.{nd}f}M"


def parse_isd(s) -> pd.Timestamp:
    """TPIT / RPG in-service strings -> Timestamp. '2027-05-15', '2029-09', '2031', 'summer 2030', 'fall 2031'. NaT if unreadable."""
    if s is None or (isinstance(s, float) and np.isnan(s)):
        return pd.NaT
    t = str(s).strip().lower()
    if not t or t in ("nan", "none", "9999-01-01"):
        return pd.NaT
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", t)
    if m:
        y = int(m.group(1))
        return pd.NaT if y >= 9000 else pd.Timestamp(y, int(m.group(2)), min(int(m.group(3)), 28))
    m = re.match(r"^(\d{4})-(\d{2})$", t)
    if m:
        return pd.Timestamp(int(m.group(1)), int(m.group(2)), 15)
    m = re.match(r"^(spring|summer|fall|autumn|winter)\s+(\d{4})$", t)
    if m:
        return pd.Timestamp(int(m.group(2)), {"spring": 4, "summer": 7, "fall": 10, "autumn": 10, "winter": 1}[m.group(1)], 15)
    m = re.match(r"^(\d{4})$", t)
    if m:
        return pd.Timestamp(int(m.group(1)), 6, 30)
    return pd.NaT


def write_labels(csv_path, spec: list[tuple[str, str, str, str]], df: pd.DataFrame) -> None:
    """Sidecar `<name>.labels.csv` with columns column,label,source,note. Every column of df must be in spec."""
    missing = [c for c in df.columns if c not in {s[0] for s in spec}]
    if missing:
        raise ValueError(f"{csv_path.name}: columns without a label: {missing}")
    bad = [s[0] for s in spec if s[1] not in LABELS]
    if bad:
        raise ValueError(f"{csv_path.name}: bad label for {bad}")
    rows = [s for s in spec if s[0] in df.columns]
    pd.DataFrame(rows, columns=["column", "label", "source", "note"]).to_csv(csv_path.with_suffix(".labels.csv"), index=False, encoding="utf-8")


# =====================================================================================================================
# Loaders (cached; the CSVs are small)
# =====================================================================================================================
@lru_cache(maxsize=1)
def load_constraints() -> pd.DataFrame:
    return pd.read_csv(CONSTRAINTS_CSV, keep_default_na=False, na_values=[""], low_memory=False)


@lru_cache(maxsize=1)
def load_planned() -> pd.DataFrame:
    return pd.read_csv(PLANNED_CSV, keep_default_na=False, na_values=[""], low_memory=False)


@lru_cache(maxsize=1)
def load_lines() -> pd.DataFrame:
    return pd.read_csv(LINES_CSV, low_memory=False)


@lru_cache(maxsize=1)
def load_substations() -> pd.DataFrame:
    s = pd.read_csv(SUBSTATIONS_CSV, low_memory=False)
    return s.set_index("sub", drop=False)


@lru_cache(maxsize=1)
def load_shift_factors() -> pd.DataFrame:
    return pd.read_csv(SHIFT_FACTORS_CSV, low_memory=False)


@dataclass(frozen=True)
class UnitCost:
    low: float
    mid: float
    high: float
    unverified: bool
    key: str


@lru_cache(maxsize=1)
def load_unit_costs() -> dict:
    """{(item, kv string): UnitCost} from data/unit_costs.csv, blank cells falling back to the mid cell (same convention as gridsite.costs)."""
    u = pd.read_csv(UNIT_COSTS_CSV, dtype=str, keep_default_na=False)
    out = {}
    for _, r in u.iterrows():
        def num(c, default):
            return float(r[c]) if r[c] != "" else default
        mid = num("mid_usd", float("nan"))
        low, high = num("low_usd", mid), num("high_usd", mid)
        mid = mid if not math.isnan(mid) else low
        low = mid if math.isnan(low) else low
        high = mid if math.isnan(high) else high
        out[(r["item"], r["kv"])] = UnitCost(low, mid, high, "UNVERIFIED" in r["note"], f"{r['item']} [{r['kv']}]")
    return out


def uc(item: str, kv: str) -> UnitCost:
    try:
        return load_unit_costs()[(item, kv)]
    except KeyError:
        raise KeyError(f"unit_costs.csv has no row ({item!r}, {kv!r})") from None


# =====================================================================================================================
# Binding episodes from the RT and DAM windows
# =====================================================================================================================
@lru_cache(maxsize=1)
def window_tables():
    """(sced, dam) shadow-price tables from derived/*.parquet, columns kept small."""
    sced = pd.read_parquet(SHADOW_SCED, columns=["sced_ts", "date_cpt", "i15", "cid", "limit", "flow", "shadow_price"])
    dam = pd.read_parquet(SHADOW_DAM, columns=["interval_start_utc", "date_cpt", "h", "cid", "limit", "shadow_price"])
    return sced, dam


def rt_window_days() -> float:
    sced, _ = window_tables()
    return sced["sced_ts"].nunique() * SCED_DT_MIN / 60.0 / 24.0


def _minutes(ts) -> np.ndarray:
    idx = pd.DatetimeIndex(pd.to_datetime(ts, utc=True))
    return np.asarray((idx - pd.Timestamp("1970-01-01", tz="UTC")) / pd.Timedelta(minutes=1), dtype=float)


def episode_hours(ts, dt_min: float, gap_min: float = EPISODE_GAP_MIN, tol_min: float = EPISODE_TOL_MIN) -> np.ndarray:
    """Binding time (hours) of each episode. `ts` are the start times of binding intervals (SCED runs of dt_min = 5 or DAM hours of 60).
    Two binding intervals belong to one episode if the time between the end of one and the start of the next is at most gap_min
    (+ tol_min of clock jitter). The hours of an episode count binding intervals only, not the gaps inside it."""
    m = np.unique(_minutes(ts))
    if m.size == 0:
        return np.zeros(0)
    new = np.diff(m) > dt_min + gap_min + tol_min
    ep = np.concatenate([[0], np.cumsum(new)])
    return np.bincount(ep) * dt_min / 60.0


def coverage_at(ep: np.ndarray, d: float) -> float:
    """Share of all binding hours covered when relief can be sustained for at most d hours from the start of each episode."""
    e = np.asarray(ep, float)
    return float(np.minimum(e, d).sum() / e.sum()) if e.size and e.sum() > 0 else float("nan")


def duration_at_coverage(ep: np.ndarray, x: float) -> float:
    """Smallest duration D (hours) with coverage_at(ep, D) >= x. Exact (piecewise linear). NaN if there are no episodes."""
    e = np.sort(np.asarray(ep, float))
    n = e.size
    if n == 0 or e.sum() <= 0:
        return float("nan")
    target = x * e.sum()
    s = 0.0
    for k in range(n):
        d = (target - s) / (n - k)
        if d <= e[k] + 1e-12:
            return float(max(d, 0.0))
        s += e[k]
    return float(e[-1])


@lru_cache(maxsize=1)
def element_class_map() -> dict:
    """constraint_name -> pooling class: ('line', 69|138|345|765), ('xfmr', None) or ('gtc', None)."""
    c = load_constraints()
    out = {}
    for name, g in c.groupby("constraint_name"):
        t = g["element_type"].iloc[0]
        kv = pd.to_numeric(g["kv"], errors="coerce").max()
        out[name] = ("line", kv_class(kv)) if t == "line" else (t, None)
    return out


@lru_cache(maxsize=1)
def all_episodes() -> dict:
    """{constraint_name: {'rt': episode hours array, 'dam': array}} for every element that binds (shadow price >= EPISODE_MIN_SHADOW)
    in the RT or DAM window. An element binds when any of its contingency constraints binds (union)."""
    c = load_constraints()
    cid2key = dict(zip(c["cid"], c["constraint_name"]))
    sced, dam = window_tables()
    rt = sced[sced["shadow_price"] >= EPISODE_MIN_SHADOW]
    dm = dam[dam["shadow_price"] >= EPISODE_MIN_SHADOW]
    out: dict = {}
    for key, g in rt.assign(k=rt["cid"].map(cid2key)).groupby("k"):
        out.setdefault(key, {})["rt"] = episode_hours(g["sced_ts"], SCED_DT_MIN)
    for key, g in dm.assign(k=dm["cid"].map(cid2key)).groupby("k"):
        out.setdefault(key, {})["dam"] = episode_hours(g["interval_start_utc"], DAM_DT_MIN)
    for v in out.values():
        v.setdefault("rt", np.zeros(0))
        v.setdefault("dam", np.zeros(0))
    return out


def pooled_episodes(cls: tuple, market: str) -> np.ndarray:
    """All episodes of the elements of one class (same market), concatenated."""
    cm = element_class_map()
    parts = [v[market] for k, v in all_episodes().items() if cm.get(k) == cls and v[market].size]
    return np.concatenate(parts) if parts else np.zeros(0)


def episode_basis_for(key: str, cls: tuple) -> tuple[np.ndarray, str]:
    """Episodes used for an element: its own RT episodes, else its own DAM episodes (each if >= MIN_BASIS_EPISODES), else the pooled class
    (RT then DAM, if >= MIN_POOL_EPISODES), else all lines pooled. Returns (episode hours, basis label)."""
    own = all_episodes().get(key, {"rt": np.zeros(0), "dam": np.zeros(0)})
    if own["rt"].size >= MIN_BASIS_EPISODES:
        return own["rt"], "own_rt"
    if own["dam"].size >= MIN_BASIS_EPISODES:
        return own["dam"], "own_dam"
    for mkt in ("rt", "dam"):
        p = pooled_episodes(cls, mkt)
        if p.size >= MIN_POOL_EPISODES:
            return p, f"pooled_{mkt}_{cls[0]}{'' if cls[1] is None else cls[1]}"
    cm = element_class_map()
    allp = np.concatenate([v["rt"] for k, v in all_episodes().items() if cm.get(k, ("",))[0] == "line" and v["rt"].size] or [np.zeros(0)])
    return allp, "pooled_rt_all_lines"


# =====================================================================================================================
# Elements: one row per monitored element (constraint_name), with facts and the selection flags
# =====================================================================================================================
def _first_valid(series):
    for v in series:
        if isinstance(v, str) and v:
            return v
    return ""


@lru_cache(maxsize=1)
def all_elements() -> pd.DataFrame:
    """Every monitored element (constraint_name) in constraints.csv with its cids and summary columns. Not yet filtered.
    The main cid is the contingency with the largest allocated IMM rent, then the largest window rent, then the best priority rank."""
    c = load_constraints().copy()
    c["kvn"] = pd.to_numeric(c["kv"], errors="coerce")
    c["_r"] = c["annual_rent_alloc_musd"].fillna(0)
    c["_w"] = c["rt_rent_est_musd"].fillna(0) + c["dam_rent_est_musd"].fillna(0)
    c = c.sort_values(["constraint_name", "_r", "_w", "priority_rank"], ascending=[True, False, False, True])
    main = c.drop_duplicates("constraint_name", keep="first").set_index("constraint_name")
    g = c.groupby("constraint_name", sort=False)
    agg = g.agg(n_cids=("cid", "size"), best_rank=("priority_rank", "min"), priority_score=("priority_score", "max"), priority_current=("priority_current", "max"),
                imm_rent_2025_musd=("annual_rent_2025_musd", "max"), imm_hours_2025=("annual_binding_hours", "max"), rt_runs=("rt_n_runs_binding", "sum"),
                dam_hours_sum=("dam_hours_binding", "sum"), window_rent_rt_musd=("rt_rent_est_musd", "sum"), window_rent_dam_musd=("dam_rent_est_musd", "sum"),
                limit_mw=("median_limit_mw", "median"), window_binding=("window_binding", "any"), kvmax=("kvn", "max"))
    cids = g["cid"].apply(list).rename("cids")

    def first_str(col):
        d = c[c[col].notna() & (c[col] != "")]
        return d.drop_duplicates("constraint_name").set_index("constraint_name")[col]
    out = agg.join(cids)
    for col in ("from_station", "to_station", "from_sub", "to_sub", "annual_element", "annual_rent_market"):
        out[col] = first_str(col).reindex(out.index).fillna("")
    out["element"] = main["element"]
    out["element_type"] = main["element_type"]
    out["main_cid"] = main["cid"]
    out["endpoint_how"] = main["endpoint_how"]
    out["endpoint_conf"] = main["endpoint_conf"]
    out = out.rename(columns={"annual_rent_market": "imm_market"})
    out["kv"] = out["kvmax"].round()
    out["kv_class"] = [None if t == "gtc" else kv_class(v) for t, v in zip(out["element_type"], out["kvmax"])]
    out["priority_score"] = out["priority_score"].astype(float)
    out = out.reset_index().rename(columns={"constraint_name": "element_key"})
    out["best_rank"] = out["best_rank"].astype(int)
    out["rt_runs"] = out["rt_runs"].astype(int)
    out["dam_hours_sum"] = out["dam_hours_sum"].astype(int)
    out["window_binding"] = out["window_binding"].astype(bool)
    return out[["element_key", "element", "element_type", "kv", "kv_class", "from_station", "to_station", "from_sub", "to_sub", "cids", "n_cids", "main_cid", "best_rank",
                "priority_score", "priority_current", "annual_element", "imm_rent_2025_musd", "imm_market", "imm_hours_2025", "rt_runs", "dam_hours_sum",
                "window_rent_rt_musd", "window_rent_dam_musd", "limit_mw", "window_binding", "endpoint_how", "endpoint_conf"]]


def select_elements() -> pd.DataFrame:
    """The elements this phase studies: the top TOP_PRIORITY_ELEMENTS by priority rank, every element with a 2025 IMM top-10 rent (even if it
    never bound in the window) and the WINDOW_TOP_ELEMENTS best window elements not yet chosen. Adds `selection` and `upgrade_id` (U01.. in rank order)."""
    E = all_elements().sort_values("best_rank").reset_index(drop=True)
    flags: dict[str, list[str]] = {}
    for k in E["element_key"].head(TOP_PRIORITY_ELEMENTS):
        flags.setdefault(k, []).append("priority_top")
    for k in E.loc[E["imm_rent_2025_musd"].notna(), "element_key"]:
        flags.setdefault(k, []).append("imm_structural")
    rest = E[~E["element_key"].isin(flags) & E["window_binding"] & (E["element_type"] != "gtc")]
    for k in rest.sort_values("priority_current", ascending=False)["element_key"].head(WINDOW_TOP_ELEMENTS):
        flags.setdefault(k, []).append("window_top")
    S = E[E["element_key"].isin(flags)].copy().reset_index(drop=True)
    S["selection"] = S["element_key"].map(lambda k: "+".join(flags[k]))
    S["upgrade_id"] = [f"U{i + 1:02d}" for i in range(len(S))]
    return S


# =====================================================================================================================
# Planned upgrades: automatic match from constraints.csv plus a short, documented review of the structural elements
# =====================================================================================================================
# ASSUMPTION (analyst review of data/planned_projects.csv on 10 Oct 2026; each entry is a project title / id actually present in that table):
# the automatic match (constraints.csv upgrade_*) is a name / bus overlap, so for the structural elements the review says what it is.
PLANNED_REVIEW = {
    "WESTEX": dict(
        planned="yes", ids="PBRP-IMPORT1;PBRP-IMPORT2;PBRP-IMPORT3", isd="2031", status="approved", conf=0.70, tsp="Oncor / LCRA TSC / AEP",
        note="Permian Basin Reliability Plan: three 765 kV import paths ($9.06B, about 1,255 miles), approved by the PUCT in Project 55718; ERCOT's planning limit for WESTEX rises "
             "12,240 -> 16,200 MW between 2027 and 2030 (data/gtc.csv). Interface-level: it relieves the interface, not a single branch."),
    "35100__A": dict(
        planned="yes", ids="TPIT-105326", isd="2028-12-15", status="planned", conf=0.90, tsp="ONCOR",
        note="Parker Switch-Hicks Switch 345 kV line reconductor, 23.0 circuit-miles (the title names both ends)."),
    "15060__B": dict(
        planned="adjacent", ids="TPIT-81305A;TPIT-81305B", isd="2026-09-30", status="planned", conf=0.30, tsp="ONCOR",
        note="Expanse-Vealmoor 138 kV rebuild (Oncor, 28.3 circuit-miles) meets the same Vealmoor 138 kV node but is a different line; it can shift flows, it is not "
             "the monitored Vealmoor-Koch Tap line. In-service date is 10 days before the as-of date and TPIT (July base) still says planned."),
    "LNGSW-CONSW LINES": dict(
        planned="yes", ids="TPIT-81268", isd="2028-05-15", status="planned", conf=0.90, tsp="ONCOR",
        note="Longshore-Consavvy 345 kV double-circuit line rebuild, 66.8 circuit-miles (the title names both ends). Also TPIT-102846 Longshore 345 kV switch rebuild "
             "(2027-05-15) and the Longshore 765 kV build-out (TPIT-99124/99126, 2028)."),
    "6520__E": dict(
        planned="likely", ids="TPIT-81418;TPIT-93466;TPIT-81440", isd="2026-12-15", status="planned", conf=0.60, tsp="ONCOR", miles=27.2,
        note="Reiter-Yarbrough 138 kV rebuild (27.2 mi, 2026-12-15), Odessa EHV-Reiter 138 kV double-circuit rebuild (2027-12-15) and Yarbrough-Pecos Trail rebuild "
             "(9.4 mi, 2027-09-15) lie on the Odessa EHV-Yarbrough path; Salt Flat Road-Barr Ranch-Reiter second circuit (TPIT-81175) went in service 2025-12-15. "
             "Whether the monitored branch is one of these segments is not verifiable from public data."),
    "LARDVN_LASCRU1_1": dict(
        planned="adjacent", ids="RPG-26RPG004;TPIT-99744;TPIT-76076", isd="2029-09", status="under_rpg_review", conf=0.30, tsp="AEP",
        note="AEPSC Laredo Area Upgrade Project ($178M, Tier 2, under RPG review, 2029-09); Laredo VFT North 138 kV breaker addition (TPIT-99744, 2026-06-11) and the "
             "Laredo VFT North-North Laredo Switch rebuild (TPIT-76076, in service 2025-08) are next to the monitored line."),
    "587__A": dict(
        planned="likely", ids="TPIT-110248;TPIT-110519;RTP2025-2025-NC34", isd="2032-12-01", status="planned", conf=0.60, tsp="ONCOR", miles=20.0,
        note="LWSVH is 'Highlands TNP' (PSS/E bus name TNHIGHLNDS1; bus 1974 in the RTP) so the monitored line is Highlands-Argyle: Argyle Switch-Highlands 138 kV reconductor, 20.0 mi "
             "(TPIT-110248, 2032-12-01), Argyle Switch-Corinth (TPIT-110519, 2031-03-01) and the RTP Highlands-Argyle-Corinth upgrade (RTP2025-2025-NC34, no date). Late in the decade."),
    "3170__B": dict(
        planned="identified", ids="RTP2025-2025-NC41", isd="", status="rtp_identified", conf=0.70,
        note="RTP 2025 'Simpson Stuart (2986) to Buntin Drive (2985) 138-kV line upgrade': bus 2985 is BNTDR; SMPST holds buses 2987 / 2988 in the current model (renumbered). "
             "Identified in the RTP with no date or cost."),
}


def _project_row(pid: str):
    p = load_planned()
    m = p[p["project_id"] == pid]
    return m.iloc[0] if len(m) else None


def _project_miles(pr) -> float:
    """Circuit-miles of a TPIT project row (rebuilt or new, whichever is given); NaN if neither."""
    if pr is None:
        return float("nan")
    v = [x for x in (pr["circuit_miles_rebuilt"], pr["circuit_miles_new"]) if x == x and x is not None]
    return float(max(v)) if v else float("nan")


def planned_for(el: pd.Series) -> dict:
    """Planned-upgrade information for one element: the best automatic match over its cids, overridden by PLANNED_REVIEW where a review exists."""
    c = load_constraints()
    g = c[c["cid"].isin(el["cids"])]
    g = g.sort_values(["upgrade_match_conf"], ascending=False, na_position="last")
    best = g.iloc[0]
    conf = best["upgrade_match_conf"]
    out = dict(planned="no", planned_id="", planned_name="", planned_isd="", planned_status="", planned_conf=np.nan, planned_method="",
               planned_tsp="", planned_miles=np.nan, planned_note="", planned_source="none")
    if isinstance(best["upgrade_project_id"], str) and np.isfinite(conf if conf == conf else np.nan):
        status = str(best["upgrade_status"])
        isd = str(best["upgrade_in_service_date"]) if isinstance(best["upgrade_in_service_date"], str) else ""
        if status == "in_service":
            cat = "done"
        elif conf >= PLANNED_YES_CONF:
            cat = "yes"
        elif conf >= UPGRADE_CONF_MIN:
            cat = "likely"
        elif conf >= PLANNED_HINT_CONF:
            cat = "hint"
        else:
            cat = "no"
        if cat in ("yes", "likely") and status == "rtp_identified" and not isd:
            cat = "identified"
        first = str(best["upgrade_project_id"]).split(";")[0]
        pr = _project_row(first)
        out.update(planned=cat, planned_id=str(best["upgrade_project_id"]), planned_name=str(best["upgrade_project_name"])[:160],
                   planned_isd=isd, planned_status=status, planned_conf=float(conf), planned_method=str(best["upgrade_match_method"]),
                   planned_tsp=("" if pr is None or not isinstance(pr["tsp"], str) else pr["tsp"]),
                   planned_miles=_project_miles(pr),
                   planned_source="auto:constraints.csv")
        if el["element_type"] == "gtc" and cat in ("yes", "likely"):
            out["planned_note"] = "interface-level: the listed projects relieve the interface limit, not a single branch"
    rev = PLANNED_REVIEW.get(el["element_key"])
    if not rev and str(el["endpoint_how"]).startswith("rejected"):
        out.update(planned="no", planned_id="", planned_name="", planned_isd="", planned_status="", planned_conf=np.nan, planned_method="", planned_tsp="", planned_miles=np.nan,
                   planned_note="the placement of this pair was rejected in phase 2 (ends too far apart), so the automatic project match is ignored", planned_source="ignored:rejected endpoints")
    if rev:
        pr = _project_row(rev["ids"].split(";")[0])
        out.update(planned=rev["planned"], planned_id=rev["ids"], planned_name=("" if pr is None else str(pr["name"])[:160]), planned_isd=rev["isd"],
                   planned_status=rev["status"], planned_conf=rev["conf"], planned_method="review",
                   planned_tsp=rev.get("tsp", "" if pr is None or not isinstance(pr["tsp"], str) else pr["tsp"]),
                   planned_miles=_project_miles(pr),
                   planned_note=rev["note"], planned_source="review:PLANNED_REVIEW")
    return out


# =====================================================================================================================
# Length, limit, binding hours of an element
# =====================================================================================================================
@lru_cache(maxsize=1)
def _line_pairs() -> dict:
    l = load_lines()
    l = l[l["from_sub"].notna() & l["to_sub"].notna() & (l["length_km"] > 0)]
    out: dict = {}
    for r in l.itertuples():
        out.setdefault(frozenset((r.from_sub, r.to_sub)), []).append((float(r.voltage_kv), float(r.length_km), str(r.source), int(r.n_circuits)))
    return out


@lru_cache(maxsize=1)
def class_median_length_km() -> dict:
    """kV class -> median OSM line length (km) among lines whose two ends are ERCOT substations (lines.csv)."""
    l = load_lines()
    l = l[(l["source"] == "osm") & l["from_sub"].notna() & l["to_sub"].notna() & (l["length_km"] > 0.2)].copy()
    l["kc"] = l["voltage_kv"].map(kv_class)
    return {int(k): float(g["length_km"].quantile(LENGTH_FALLBACK_QUANTILE)) for k, g in l.groupby("kc") if k in (69, 138, 345)}


def element_length(el: pd.Series, planned: dict) -> dict:
    """Route length of a line element: lines.csv (same pair, same voltage class) > TPIT circuit-miles of a reviewed / strongly matched project >
    station distance x ROUTE_FACTOR (both ends placed) > class median. Returns km, miles and a basis label."""
    out = dict(length_km=np.nan, length_mi=np.nan, length_basis="n/a")
    if el["element_type"] != "line":
        return out
    a, b = el["from_sub"], el["to_sub"]
    kc = _kc(el["kv_class"])
    if str(el["endpoint_how"]).startswith("rejected"):
        a = b = ""                    # phase 2 rejected the placement of this pair (ends too far apart): do not trust the station names
    if a and b and a != b:
        recs = [r for r in _line_pairs().get(frozenset((a, b)), []) if kv_class(r[0]) == kc]
        if recs:
            km = float(np.median([r[1] for r in recs]))
            srcs = sorted({r[2] for r in recs})
            return dict(length_km=km, length_mi=km / KM_PER_MILE, length_basis=f"lines.csv:{'+'.join(srcs)}({len(recs)})")
    rv = PLANNED_REVIEW.get(el["element_key"], {})
    if rv.get("miles"):
        return dict(length_km=rv["miles"] * KM_PER_MILE, length_mi=float(rv["miles"]), length_basis=f"review: {rv['ids'].split(';')[0]} circuit-miles of the same corridor")
    pm = planned.get("planned_miles", np.nan)
    if np.isfinite(pm) and pm > 0 and planned.get("planned") in ("yes",) and planned.get("planned_conf", 0) >= PLANNED_YES_CONF:
        return dict(length_km=pm * KM_PER_MILE, length_mi=float(pm), length_basis=f"planned:{planned['planned_id'].split(';')[0]} circuit-miles")
    if a and b and a != b:
        s = load_substations()
        if a in s.index and b in s.index:
            ra, rb = s.loc[a], s.loc[b]
            if (ra["match_conf"] >= ANCHOR_CONF and rb["match_conf"] >= ANCHOR_CONF and np.isfinite(ra["lat"]) and np.isfinite(rb["lat"])):
                km = haversine_km(ra["lat"], ra["lon"], rb["lat"], rb["lon"]) * ROUTE_FACTOR
                return dict(length_km=km, length_mi=km / KM_PER_MILE, length_basis=f"station_distance x {ROUTE_FACTOR:g}")
    cm = class_median_length_km().get(kc)
    if cm:
        return dict(length_km=cm, length_mi=cm / KM_PER_MILE, length_basis=f"class_median {kc} kV (lines.csv); ends not placed")
    return out


@lru_cache(maxsize=1)
def class_median_limit_mw() -> dict:
    """(element_type, kv class or None) -> median of the element limits over window-binding elements."""
    E = all_elements()
    E = E[E["limit_mw"].notna() & (E["element_type"] != "gtc")]
    out = {}
    for (t, kc), g in E.groupby(["element_type", "kv_class"], dropna=False):
        out[(t, None if pd.isna(kc) else int(kc))] = float(g["limit_mw"].quantile(LIMIT_FALLBACK_QUANTILE))
    return out


def element_limit(el: pd.Series) -> dict:
    if np.isfinite(el["limit_mw"]):
        return dict(limit_mw=float(el["limit_mw"]), limit_basis="window median (SCED / DAM limit)")
    kc = None if el["kv_class"] is None or pd.isna(el["kv_class"]) else int(el["kv_class"])
    v = class_median_limit_mw().get((el["element_type"], kc))
    if v:
        return dict(limit_mw=v, limit_basis=f"class median {el['element_type']} {kc} kV (window); element never priced in the window")
    return dict(limit_mw=np.nan, limit_basis="unknown")


def limit_swing(el: pd.Series) -> float:
    """(q95 - q05) / median of the SCED limit over the element's listed rows: a visible swing means ratings already follow ambient temperature."""
    sced, _ = window_tables()
    x = sced.loc[sced["cid"].isin(el["cids"]) & (sced["limit"] > 0), "limit"]
    if len(x) < SWING_MIN_ROWS:
        return float("nan")
    return float((x.quantile(SWING_QUANTILES[1]) - x.quantile(SWING_QUANTILES[0])) / x.median())


@lru_cache(maxsize=1)
def imm_hours_per_musd() -> tuple[float, int]:
    """Median IMM binding hours per $M of 2025 RT rent over the elements that have both (DERIVED); used where only the rent is known."""
    E = all_elements()
    x = E[E["imm_hours_2025"].notna() & E["imm_rent_2025_musd"].notna() & (E["imm_market"] == "RT")]
    r = (x["imm_hours_2025"] / x["imm_rent_2025_musd"])
    return float(r.median()), int(len(r))


def annual_hours(el: pd.Series, eps: dict) -> dict:
    """Annual binding hours: IMM 2025 hours (REAL) > IMM rent x median hours per $M (DERIVED) > window hours x 365 / window days x persistence."""
    if np.isfinite(el["imm_hours_2025"]):
        return dict(annual_hours=float(el["imm_hours_2025"]), annual_hours_basis=f"IMM 2025 ({el['imm_market'] or 'RT'})")
    if np.isfinite(el["imm_rent_2025_musd"]):
        k, n = imm_hours_per_musd()
        return dict(annual_hours=float(el["imm_rent_2025_musd"] * k), annual_hours_basis=f"IMM rent x {k:.1f} h per $M (median of {n} elements)")
    rt_h, dam_h = eps["rt"].sum(), eps["dam"].sum()
    if rt_h > 0:
        return dict(annual_hours=float(rt_h * 365.0 / rt_window_days() * WINDOW_PERSISTENCE), annual_hours_basis=f"RT window x 365 / {rt_window_days():.1f} d x {WINDOW_PERSISTENCE}")
    if dam_h > 0:
        return dict(annual_hours=float(dam_h * 365.0 / DAM_WINDOW_DAYS * WINDOW_PERSISTENCE), annual_hours_basis=f"DAM window x 365 / {DAM_WINDOW_DAYS:g} d x {WINDOW_PERSISTENCE}")
    return dict(annual_hours=np.nan, annual_hours_basis="no data")


# =====================================================================================================================
# Lead times from TPIT: months from the RPG submittal date to the (projected or actual) in-service date, by project class
# =====================================================================================================================
def _tpit_class(title, mva, mi_new, mi_reb) -> str:
    t = str(title).lower()
    if (mva or 0) > 0 or "autotransformer" in t:
        return "autotransformer"
    if (mi_new or 0) > 0.5:
        return "new_line"
    if (mi_reb or 0) > 0.5:
        return "line_rebuild_reconductor"
    return "substation_other"


@lru_cache(maxsize=1)
def tpit_lead_table() -> pd.DataFrame:
    """Quartiles of (in service - RPG submittal) in months for TPIT projects, by basis (pipeline = projected in-service after the as-of date,
    completed = actual in-service), class (line_rebuild_reconductor / new_line / autotransformer / substation_other) and kV class (69 / 138 / 345 / all)."""
    import openpyxl
    wb = openpyxl.load_workbook(TPIT_XLSX, read_only=True, data_only=True)
    frames = []
    for prefix in ("Completed", "Planned", "Future"):
        name = next(n for n in wb.sheetnames if n.startswith(prefix))
        rows = list(wb[name].iter_rows(min_row=2, values_only=True))
        df = pd.DataFrame(rows[1:], columns=[str(h) for h in rows[0]])
        df["sheet"] = prefix
        frames.append(df)
    t = pd.concat(frames, ignore_index=True)

    def col(prefix):
        return next(c for c in t.columns if c.startswith(prefix))
    x = pd.DataFrame({
        "sheet": t["sheet"], "title": t[col("Project Title")],
        "kv": pd.to_numeric(t[col("Service Level kV")], errors="coerce"),
        "mi_new": pd.to_numeric(t[col("Trans Circuit Miles New")], errors="coerce"),
        "mi_reb": pd.to_numeric(t[col("Trans Circuit Miles Rebuilt")], errors="coerce"),
        "mva": pd.to_numeric(t[col("Autotransformer Capacity")], errors="coerce"),
        "d_sub": pd.to_datetime(t[col("Date Submitted TO ERCOT")], errors="coerce"),
        "isd_proj": pd.to_datetime(t[col("Projected In-Service")], errors="coerce"),
        "isd_act": pd.to_datetime(t[col("Actual In-Service")], errors="coerce")})
    x["cls"] = [_tpit_class(a, b, c, d) for a, b, c, d in zip(x["title"], x["mva"], x["mi_new"], x["mi_reb"])]
    x["kvc"] = x["kv"].map(lambda v: str(kv_class(v)))
    comp = x[(x["sheet"] == "Completed") & x["d_sub"].notna() & x["isd_act"].notna()].assign(m=lambda d: (d["isd_act"] - d["d_sub"]).dt.days / 30.4375)
    pipe = x[(x["sheet"] != "Completed") & x["d_sub"].notna() & (x["isd_proj"] > AS_OF)].assign(m=lambda d: (d["isd_proj"] - d["d_sub"]).dt.days / 30.4375)
    out = []
    for basis, d in (("pipeline", pipe), ("completed", comp)):
        d = d[d["m"] > 0]
        groups = list(d.groupby(["cls", "kvc"])) + [((c, "all"), g) for c, g in d.groupby("cls")]
        for (cls, kvc), g in groups:
            out.append(dict(basis=basis, cls=cls, kvc=kvc, n=len(g), q25=g["m"].quantile(0.25), med=g["m"].median(), q75=g["m"].quantile(0.75)))
    return pd.DataFrame(out)


OPTION_TPIT_CLASS = {"reconductor": "line_rebuild_reconductor", "rebuild": "line_rebuild_reconductor", "new_parallel_line": "new_line",
                     "autotransformer": "autotransformer", "substation_rebuild": "substation_other"}


def lead_months(option: str, kc: int | None) -> dict:
    """(low, mid, high) months for a wires option: TPIT pipeline quartiles for its class (own kV class if >= LEAD_MIN_N rows, else all kV);
    an autotransformer is floored by the large-power-transformer lead time plus installation; DLR / AAR and the 765 kV plan use their own rules."""
    if option == "dlr_aar":
        lo, mi, hi = DLR_LEAD_MONTHS
        return dict(lead_low=lo, lead_mid=mi, lead_high=hi, lead_basis="ASSUMPTION: sensors + software; a ratings change is not an RPG project (NP 3.11.4.1(4))")
    if option == "plan_765kv":
        lo, mi, hi = (max((pd.Timestamp(d) - AS_OF).days / 30.4375, 0.0) for d in PLAN_765_ISD)
        return dict(lead_low=lo, lead_mid=mi, lead_high=hi, lead_basis="REAL: PBRP 765 kV import paths in service 2031 (C&N Dec 2025 p.8), months from the as-of date")
    cls = OPTION_TPIT_CLASS[option]
    t = tpit_lead_table()
    t = t[(t["basis"] == "pipeline") & (t["cls"] == cls)]
    r = t[t["kvc"] == str(kc)] if kc else t.iloc[0:0]
    scope = f"{kc} kV"
    if not len(r) or r["n"].iloc[0] < LEAD_MIN_N:
        r, scope = t[t["kvc"] == "all"], "all kV"
    q25, med, q75, n = (float(r["q25"].iloc[0]), float(r["med"].iloc[0]), float(r["q75"].iloc[0]), int(r["n"].iloc[0]))
    basis = f"DERIVED: TPIT {cls} {scope} pipeline, RPG submittal to projected in service, median and quartiles, n={n}"
    done = tpit_lead_table()
    done = done[(done["basis"] == "completed") & (done["cls"] == cls)]
    dr = done[done["kvc"] == str(kc)] if kc else done.iloc[0:0]
    dscope = f"{kc} kV"
    if not len(dr) or dr["n"].iloc[0] < LEAD_MIN_N:
        dr, dscope = done[done["kvc"] == "all"], "all kV"
    if len(dr) and dr["n"].iloc[0] >= LEAD_MIN_N and float(dr["med"].iloc[0]) < q25:
        q25 = float(dr["med"].iloc[0])
        basis += f"; low = median of {int(dr['n'].iloc[0])} completed {dscope} projects (actual)"
    if option == "autotransformer":
        lpt = [m + XFMR_INSTALL_MONTHS for m in LPT_LEAD_MONTHS]
        q25, med, q75 = max(q25, lpt[0]), max(med, lpt[1]), max(q75, lpt[2])
        basis += (f"; floored by large-power-transformer lead {LPT_LEAD_MONTHS[0]:g}/{LPT_LEAD_MONTHS[1]:g}/{LPT_LEAD_MONTHS[2]:g} months "
                  f"+ {XFMR_INSTALL_MONTHS:g} install (DOE 2024, NLR 2026)")
    return dict(lead_low=q25, lead_mid=med, lead_high=q75, lead_basis=basis)


# =====================================================================================================================
# Wires options and their costs (data/unit_costs.csv x quantity)
# =====================================================================================================================
def _band(lo: float, mid: float, hi: float) -> tuple[float, float, float]:
    return min(lo, mid * (1 - ESTIMATE_BAND_LOW)), mid, max(hi, mid * (1 + ESTIMATE_BAND_HIGH))


def _tri(u: UnitCost, qty: float) -> tuple[float, float, float]:
    return u.low * qty, u.mid * qty, u.high * qty


def _sum(*ts) -> tuple[float, float, float]:
    return tuple(sum(t[i] for t in ts) for i in range(3))


def xfmr_high_kv(el: pd.Series) -> int:
    """High-side voltage class of a transformer element: the kV_max of its station when known (constraints.csv lists only one end's kV)."""
    s = load_substations()
    sub = el["from_sub"] or el["to_sub"]
    if sub and sub in s.index and np.isfinite(s.loc[sub, "kv_max"]):
        return kv_class(s.loc[sub, "kv_max"]) or 138
    return _kc(el["kv_class"]) or 138


def wires_options(el: pd.Series, length_mi: float, limit_mw: float, swing: float) -> list[dict]:
    """All plausible wires options for an element with a low / mid / high cost in $M (2024$), the formula, rating uplift and lead time."""
    opts: list[dict] = []

    def add(option, label, tri, formula, uplift, ccn, process, uses_unv, kc_for_lead):
        tri = _band(*tri)
        d = dict(upgrade_id=el["upgrade_id"], option=option, label=label, low_musd=tri[0] / 1e6, mid_musd=tri[1] / 1e6, high_musd=tri[2] / 1e6,
                 formula=formula, uplift=uplift, ccn_needed=ccn, process=process, uses_unverified_cell=bool(uses_unv))
        d.update(lead_months(option, kc_for_lead))
        opts.append(d)

    et, kc = el["element_type"], _kc(el["kv_class"])
    if et == "line" and np.isfinite(length_mi):
        k = kc if kc in (69, 138, 345) else 138
        bay = uc("substation: add one bay (ring bus)", "138" if k <= 138 else "345")
        term = tuple(2 * TERMINAL_BAY_FRACTION * v for v in (bay.low, bay.mid, bay.high))
        term_txt = f"2 line ends x {TERMINAL_BAY_FRACTION:g} x {fm(bay.mid / 1e6)} bay"
        rec = uc("reconductor per circuit", "345" if k == 345 else "138")
        add("reconductor", "reconductor the monitored circuit", _sum(_tri(rec, length_mi), term),
            f"{length_mi:.1f} mi x {fm(rec.mid / 1e6, 2)}/mi ({rec.key}{'; 69 kV priced at the 138 kV row' if k == 69 else ''}) + {term_txt}",
            UPLIFT["reconductor"], False, "Tier 4 neutral if ratings only (NP 3.11.4.3(1)(f)(viii)); a 345 kV reconductor over 1 mile is Tier 3", rec.unverified, k)
        if k in (69, 138):
            rb = uc("rebuild of existing single-circuit line", str(k))
            rb_txt = rb.key
        else:
            n345 = uc("new single-circuit line", "345")
            rb = UnitCost(n345.low * REBUILD_345_FROM_NEW, n345.mid * REBUILD_345_FROM_NEW, n345.high * REBUILD_345_FROM_NEW, n345.unverified, n345.key)
            rb_txt = f"{n345.key} x {REBUILD_345_FROM_NEW:.2f} (138 kV rebuild / new ratio)"
        add("rebuild", "rebuild on the existing route", _sum(_tri(rb, length_mi), term),
            f"{length_mi:.1f} mi x {fm(rb.mid / 1e6, 2)}/mi ({rb_txt}) + {term_txt}", UPLIFT["rebuild"], False,
            "Tier 3 or 4; no CCN if no new right of way (16 TAC 25.101(c)(5)(B))", rb.unverified, k)
        nl = uc("new single-circuit line", str(k))
        bays = tuple(NEW_LINE_BAYS * v for v in (bay.low, bay.mid, bay.high))
        add("new_parallel_line", "new parallel single-circuit line", _sum(_tri(nl, length_mi), bays),
            f"{length_mi:.1f} mi x {fm(nl.mid / 1e6, 2)}/mi ({nl.key}) + {NEW_LINE_BAYS} x {fm(bay.mid / 1e6)} bay", UPLIFT["new_parallel_line"], True,
            "Tier 2 (CCN, under $200M) or Tier 1 ($200M or more); CCN from the PUCT", nl.unverified or bay.unverified, k)
        dl = uc("dynamic line rating, sensors and software", "69-345")
        floor = DLR_MIN_COST_MUSD * 1e6
        tri = tuple(max(v, floor) for v in _tri(dl, length_mi))
        followed = bool(np.isfinite(swing) and swing >= AMBIENT_SWING_MIN)
        up = UPLIFT["dlr_aar_after_aar"] if followed else UPLIFT["dlr_aar_static"]
        add("dlr_aar", "dynamic / ambient-adjusted line rating", tri,
            f"{length_mi:.1f} mi x {fm(dl.mid / 1e6, 3)}/mi ({dl.key}), floor {fm(DLR_MIN_COST_MUSD, 2)}; uplift {up:.0%} "
            f"({'SCED limit already follows ambient temperature, swing ' + format(swing, '.0%') if followed else 'static limit in the window'})",
            up, False, "a ratings change is not an RPG project (NP 3.11.4.1(4))", dl.unverified, k)
    if et == "xfmr":
        hv = xfmr_high_kv(el)
        mva = next((s for s in XFMR_STEPS_MVA if s >= (limit_mw if np.isfinite(limit_mw) else 0)), XFMR_STEPS_MVA[-1])
        if hv >= 345:
            rate = uc("power transformer (autotransformer) 345/138 kV", "345/138")
            bays = _sum(_tri(uc("substation: add one bay (ring bus)", "345"), 1), _tri(uc("substation: add one bay (ring bus)", "138"), 1))
            bay_txt = "1 x 345 kV bay + 1 x 138 kV bay"
        else:
            rate = UnitCost(XFMR_138_69_USD_PER_MVA, XFMR_138_69_USD_PER_MVA, XFMR_138_69_USD_PER_MVA, True, "MISO MTEP24 Table 2.3-7, 138/69 kV (layout inferred)")
            bays = _tri(uc("substation: add one bay (ring bus)", "138"), 2)
            bay_txt = "2 x 138 kV bay"
        add("autotransformer", f"parallel {'345/138' if hv >= 345 else '138/69'} kV autotransformer, {mva} MVA", _sum(_tri(rate, mva), bays),
            f"{mva} MVA x ${rate.mid:,.0f}/MVA ({rate.key}) + {bay_txt}; size = smallest standard step >= the {limit_mw:.0f} MVA rating",
            UPLIFT["autotransformer"], False, "345/138 kV autotransformer capacity is Tier 3 below $50M (NP 3.11.4.3(1)(c)(ii)); substations need no CCN",
            rate.unverified, 345 if hv >= 345 else 138)
    if et != "gtc":
        k = kc if kc in (69, 138, 345) else 138
        ss = uc("substation: new 4-position substation (ring bus)", str(SUBSTATION_REBUILD_KV_ROW[k]))
        add("substation_rebuild", "substation rebuild / new 4-position ring bus at the limiting station", _tri(ss, 1),
            f"1 x {fm(ss.mid / 1e6)} ({ss.key}); helps only if the limit is substation equipment, which public data cannot show", UPLIFT["substation_rebuild"], False,
            "Tier 3 or 4; substations need no CCN (16 TAC 25.101(c)(2))", ss.unverified, k)
    if et == "gtc" and el["element_key"] == "WESTEX":
        pl = uc("Permian Basin 765-kV import paths (3 paths, about 1,255 miles)", "765")
        add("plan_765kv", "Permian Basin Reliability Plan: three 765 kV import paths (approved)", _tri(pl, 1),
            f"3 paths, 1,255 miles, {fm(pl.mid / 1e6, 0)} in total (ERCOT PBRP study; PUCT Project 55718); it serves load growth as well as WESTEX",
            UPLIFT["plan_765kv"], True, "approved by the PUCT (Project 55718); each TSP files its CCNs", pl.unverified, None)
    if et == "gtc" and not opts:
        opts.append(dict(upgrade_id=el["upgrade_id"], option="plan_unpriced", label="projects named in data/gtc.csv; no cost in our tables", low_musd=np.nan, mid_musd=np.nan,
                         high_musd=np.nan, formula="cost of the San Antonio South / Valley projects is not in data/planned_projects.csv", uplift=np.nan, ccn_needed=True,
                         process="ERCOT planning (RTP) and PUCT", uses_unverified_cell=False, lead_low=np.nan, lead_mid=np.nan, lead_high=np.nan, lead_basis="n/a"))
    return opts


def primary_option(el: pd.Series, opts: list[dict], planned: dict | None = None) -> dict | None:
    """The structural fix a TSP would most likely file: reconductor for a line (a rebuild if the matched or reviewed planned project is a rebuild), a parallel
    autotransformer for a transformer, the 765 kV plan for WESTEX."""
    want = {"line": "reconductor", "xfmr": "autotransformer", "gtc": "plan_765kv"}[el["element_type"]]
    if (el["element_type"] == "line" and planned and planned.get("planned") in ("yes", "likely")
            and re.search(r"rebuild", str(planned.get("planned_name", "")), flags=re.I)):
        want = "rebuild"
    for o in opts:
        if o["option"] == want:
            return o
    return opts[0] if opts else None


def cheapest_eligible(opts: list[dict], relief_frac: float) -> dict | None:
    """Cheapest option (by mid cost) whose typical rating uplift covers the relief target."""
    ok = [o for o in opts if np.isfinite(o["mid_musd"]) and np.isfinite(o["uplift"]) and o["uplift"] >= relief_frac]
    return min(ok, key=lambda o: o["mid_musd"]) if ok else None


# =====================================================================================================================
# Battery cost (NREL 2025 update, 2024$)
# =====================================================================================================================
def battery_usd_per_kw(hours: float, mult: float) -> float:
    return mult * (BATT_ENERGY_USD_KWH * hours + BATT_POWER_USD_KW)


def battery_capex_musd(mw: float, hours: float, mult: float) -> float:
    return mw * 1000.0 * battery_usd_per_kw(hours, mult) / 1e6


def battery_annual_musd(capex_musd: float) -> float:
    return capex_musd * (crf(DISCOUNT_RATE, BATT_LIFE_YR) + BATT_FOM_FRAC)


def wires_annual_musd(capex_musd: float) -> float:
    return capex_musd * WIRES_FIXED_CHARGE_RATE


# =====================================================================================================================
# One study record per selected element
# =====================================================================================================================
@lru_cache(maxsize=1)
def element_study() -> list[dict]:
    """For each selected element: planned info, length, limit, ambient swing, episodes, annual hours, wires options and the primary option."""
    out = []
    for _, el in select_elements().iterrows():
        pl = planned_for(el)
        ln = element_length(el, pl)
        lim = element_limit(el)
        sw = limit_swing(el)
        cls = element_class_map().get(el["element_key"], ("line", el["kv_class"]))
        ep, eb = episode_basis_for(el["element_key"], cls)
        own = all_episodes().get(el["element_key"], {"rt": np.zeros(0), "dam": np.zeros(0)})
        ah = annual_hours(el, own)
        opts = wires_options(el, ln["length_mi"], lim["limit_mw"], sw)
        out.append(dict(el=el, planned=pl, length=ln, limit=lim, swing=sw, episodes=ep, episode_basis=eb, own=own, annual=ah,
                        options=opts, primary=primary_option(el, opts, pl)))
    return out


# =====================================================================================================================
# Shift-factor sites (substation level) for the non-wires sizing
# =====================================================================================================================
def _tokens(q) -> set[str]:
    return set(str(q).split("+"))


def site_cap_mw(kv_max) -> float:
    kc = kv_class(kv_max)
    return SITE_MAX_MW[69] if kc == 69 else SITE_MAX_MW[138] if kc in (None, 138) else SITE_MAX_MW[345]


@lru_cache(maxsize=None)
def usable_sites(cid: str) -> pd.DataFrame | None:
    """Substation SFs of one constraint that may size a battery: rt_reg first, else DAM (divided by its shrink factor if attenuated and the shrink is
    at least UNSHRINK_FLOOR); never proxy; no weak / implausible rows; 69 kV and above; |SF| >= SF_SITE_MIN, capped at SF_NWA_CAP. None if none."""
    sf = load_shift_factors()
    x = sf[sf["cid"] == cid]
    if x.empty:
        return None
    meth = x["method"].iloc[0]
    if meth == "dcpf_proxy":
        return None
    c = load_constraints().set_index("cid")
    shrink = float(c.at[cid, "sf_dam_shrink"]) if meth == "dam_reg" and pd.notna(c.at[cid, "sf_dam_shrink"]) else 1.0
    x = x.copy()
    if meth == "dam_reg" and shrink < SHRINK_MIN:
        if shrink >= UNSHRINK_FLOOR:
            x["adj"], basis = x["sf"] / shrink, "dam_reg_unshrunk"
        else:
            return None
    else:
        x["adj"], basis = x["sf"], ("rt_reg" if meth == "rt_reg" else "dam_reg")
    bad = set(BAD_SF_QUALITY)
    x = x[x["quality"].map(lambda q: not (_tokens(q) & bad)) & x["kv_max"].fillna(0).ge(69)]
    x = x.assign(abs_sf=x["adj"].abs().clip(upper=SF_NWA_CAP))
    x = x[x["abs_sf"] >= SF_SITE_MIN]
    if x.empty:
        return None
    x["side"] = np.where(x["adj"] < 0, "import", "export")
    x["basis"] = basis
    x["cap_mw"] = x["kv_max"].map(site_cap_mw)
    x = x.sort_values(["abs_sf", "sub"], ascending=[False, True])
    return x[["sub", "sf", "adj", "abs_sf", "side", "kv_max", "zone", "match_conf", "method", "quality", "basis", "cap_mw"]].reset_index(drop=True)


def element_site_data(el: pd.Series):
    """(cid used, sites table) for an element: the main cid first, then its other cids; (None, None) if no regression SF exists."""
    for cid in [el["main_cid"]] + [c for c in el["cids"] if c != el["main_cid"]]:
        df = usable_sites(cid)
        if df is not None and len(df):
            return cid, df
    return None, None


@lru_cache(maxsize=None)
def fallback_sf(kc: int) -> tuple[float, float, int]:
    """(import, export, n) best-site |SF| for a line of kV class kc that has no regression SF: FALLBACK_SF_QUANTILE over the window-fitted RT lines of that class."""
    sf = load_shift_factors()
    cons = load_constraints().set_index("cid")
    bad = set(BAD_SF_QUALITY)
    imp, exp = [], []
    for cid, g in sf[sf["method"] == "rt_reg"].groupby("cid"):
        r = cons.loc[cid]
        if r["element_type"] != "line" or kv_class(r["kv"]) != kc:
            continue
        g = g[g["quality"].map(lambda q: not (_tokens(q) & bad)) & g["kv_max"].fillna(0).ge(69)]
        a, b = (-g.loc[g["sf"] < 0, "sf"]).clip(upper=SF_NWA_CAP), g.loc[g["sf"] > 0, "sf"].clip(upper=SF_NWA_CAP)
        if len(a):
            imp.append(float(a.max()))
        if len(b):
            exp.append(float(b.max()))
    q = FALLBACK_SF_QUANTILE
    return (float(np.quantile(imp, q)) if imp else np.nan, float(np.quantile(exp, q)) if exp else np.nan, len(imp))


def size_for_relief(relief_mw: float, sites: pd.DataFrame) -> dict:
    """Greedy MW to deliver `relief_mw` of flow relief from sites sorted by |SF|: each site takes min(cap, remaining / |SF|)."""
    rem, mw, n = float(relief_mw), 0.0, 0
    for r in sites.itertuples():
        take = min(r.cap_mw, rem / r.abs_sf)
        mw += take
        rem -= take * r.abs_sf
        n += 1
        if rem <= 1e-9:
            break
    best = sites.iloc[0]
    return dict(mw_required=(mw if rem <= 1e-6 else float("nan")), n_sites=n, shortfall_mw=max(rem, 0.0), mw_best_site=relief_mw / best["abs_sf"],
                best_site=best["sub"], best_site_kv=best["kv_max"], best_site_sf=float(best["adj"]), best_site_zone=best["zone"], basis=best["basis"])


# =====================================================================================================================
# Shadow-price-hours, zone means, value of a site: the first-order rerank
# =====================================================================================================================
@lru_cache(maxsize=1)
def sph_table() -> pd.DataFrame:
    """Annual shadow-price-hours per constraint, $/MW-yr (sum over a year of the shadow price in $/MWh x 1 h). Method per row:
       imm_rent_alloc   = IMM 2025 rent allocated to the cid / the element's median limit (REAL rent, DERIVED);
       window_rt / window_dam = window rent $/MW x 365 / window days x WINDOW_PERSISTENCE, with the elements' sum capped at the 10th-ranked IMM rent / limit
                          (an element outside the IMM top 10 earns less than that)."""
    c = load_constraints().copy()
    lim = c["median_limit_mw"]
    alloc = c["annual_rent_alloc_musd"].fillna(0)
    sph = pd.Series(np.nan, index=c.index)
    basis = pd.Series("none", index=c.index)
    has_imm = (alloc > 0) & (lim > 0)
    sph[has_imm] = alloc[has_imm] * 1e6 / lim[has_imm]
    basis[has_imm] = "imm_rent_alloc"
    rt_days = rt_window_days()
    win_rt = c["rt_rent_proxy"] * 365.0 / rt_days * WINDOW_PERSISTENCE
    win_dam = c["dam_rent_proxy"] * 365.0 / DAM_WINDOW_DAYS * WINDOW_PERSISTENCE
    use_rt = (c["rt_n_runs_binding"] >= MIN_RT_RUNS_FOR_SPH)
    use_dam = ~use_rt & (c["dam_hours_binding"] > 0)
    raw = pd.Series(np.nan, index=c.index)
    raw[use_rt] = win_rt[use_rt]
    raw[use_dam] = win_dam[use_dam]
    todo = ~has_imm & raw.notna() & (lim > 0)
    # element-level cap: sum over an element's window cids must not exceed IMM_FLOOR_RANK_RENT_MUSD / limit
    cap = IMM_FLOOR_RANK_RENT_MUSD * 1e6 / c.groupby("constraint_name")["median_limit_mw"].transform("median")
    tot = raw.where(todo).groupby(c["constraint_name"]).transform("sum")
    scale = np.minimum(1.0, cap / tot).where(todo)
    sph[todo] = (raw * scale)[todo]
    basis[todo & use_rt] = "window_rt"
    basis[todo & use_dam] = "window_dam"
    basis[todo & (scale < 1)] = basis[todo & (scale < 1)] + "_capped"
    from . import rev2                                                  # R2: the elements with an IMM 2026 year-to-date rent take their annual target from it (the same rule as the lanes)

    return rev2.anchor_sph_table(pd.DataFrame({"cid": c["cid"], "sph_usd_mw_yr": sph, "sph_basis": basis, "limit_mw": lim}))


@lru_cache(maxsize=1)
def zone8_of_sub() -> dict:
    """Substation -> the arena's load zone: the NOIE zone (LZ_AEN / CPS / LCRA / RAYBN) if it has one, else the NP4-160 settlement load zone."""
    s = load_substations()
    return {sub: (r["noie_zone"] if isinstance(r["noie_zone"], str) else (r["zone"] if isinstance(r["zone"], str) else "")) for sub, r in s.iterrows()}


@lru_cache(maxsize=1)
def _bus_maps() -> tuple[dict, dict, dict]:
    sp = pd.read_csv(SETTLEMENT_CSV, dtype=str, low_memory=False)
    z = sp["NOIE_ZONE"].where(sp["NOIE_ZONE"].notna(), sp["SETTLEMENT_LOAD_ZONE"]).fillna("")
    kv = pd.to_numeric(sp["VOLTAGE_LEVEL"], errors="coerce")
    bus = sp["ELECTRICAL_BUS"].astype(str)
    n_by_zone = z[kv >= 60].value_counts().to_dict()
    return dict(zip(bus, z)), dict(zip(bus, kv)), n_by_zone


def _allowed_cids() -> pd.Series:
    """cid -> method for the constraints whose SF magnitudes may enter the rerank: rt_reg, or dam_reg that is not attenuated (shrink >= SHRINK_MIN)."""
    sf = load_shift_factors()
    c = load_constraints().set_index("cid")
    m = sf.groupby("cid")["method"].first()
    ok = {}
    for cid, meth in m.items():
        if meth == "rt_reg":
            ok[cid] = meth
        elif meth == "dam_reg" and pd.notna(c.at[cid, "sf_dam_shrink"]) and float(c.at[cid, "sf_dam_shrink"]) >= SHRINK_MIN:
            ok[cid] = meth
    return pd.Series(ok)


@lru_cache(maxsize=1)
def zone_mean_sf() -> pd.DataFrame:
    """Mean SF over the buses of 60 kV and above of each arena load zone (absent = 0; no load by bus is public, so unweighted) for the allowed constraints."""
    allowed = _allowed_cids()
    df = pd.read_parquet(SF_LONG, columns=["bus", "cid", "sf", "method"])
    df = df[df["cid"].isin(list(allowed.index))]
    for col in ("bus", "cid", "method"):
        df[col] = df[col].astype(str)
    df = df[df["method"] == df["cid"].map(allowed)]
    bz, bkv, nz = _bus_maps()
    df["zone8"] = df["bus"].map(bz)
    df["kv"] = df["bus"].map(bkv)
    df = df[df["zone8"].notna() & (df["kv"] >= 60)]
    g = df.groupby(["zone8", "cid"])["sf"].sum().reset_index()
    g["n_buses"] = g["zone8"].map(nz)
    g["zmean"] = g["sf"] / g["n_buses"]
    return g[["zone8", "cid", "zmean"]]


@lru_cache(maxsize=1)
def value_table() -> pd.DataFrame:
    """First-order congestion value of a battery at substation n for constraint k, $/kW-yr:
         V[n,k] = min(|SF[n,k] - zone mean SF[k]|, SF_NWA_CAP) x shadow-price-hours[k] / 1000 x CAPTURE_FACTOR(spec).
       The sign of SF - zone mean tells the side: negative = import side (the node price rises when k binds: sell there), positive = export side (cheap to charge)."""
    sf = load_shift_factors()
    allowed = _allowed_cids()
    bad = set(BAD_SF_QUALITY)
    x = sf[sf["cid"].isin(list(allowed.index))].copy()
    x = x[x["method"] == x["cid"].map(allowed)]
    x = x[x["quality"].map(lambda q: not (_tokens(q) & bad)) & x["kv_max"].fillna(0).ge(69)]
    x["zone8"] = x["sub"].map(zone8_of_sub()).fillna("")
    zm = zone_mean_sf()
    x = x.merge(zm, on=["zone8", "cid"], how="left")
    x["zmean"] = x["zmean"].fillna(0.0)
    x["dsf"] = x["sf"] - x["zmean"]
    x["abs_dsf"] = x["dsf"].abs().clip(upper=SF_NWA_CAP)
    x = x.merge(sph_table()[["cid", "sph_usd_mw_yr", "sph_basis"]], on="cid", how="left")
    x = x[x["sph_usd_mw_yr"].notna() & (x["sph_usd_mw_yr"] > 0)]
    for spec, cf in CAPTURE_FACTOR.items():
        x[f"v_{spec}"] = x["abs_dsf"] * x["sph_usd_mw_yr"] / 1000.0 * cf
    x["side"] = np.where(x["dsf"] < 0, "import", "export")
    return x[["sub", "cid", "sf", "zmean", "dsf", "side", "zone8", "kv_max", "match_conf", "method", "quality", "sph_usd_mw_yr", "sph_basis", "v_2h", "v_4h"]].reset_index(drop=True)


def ref_value(cid: str, sub: str, spec: str) -> float:
    vt = value_table()
    m = vt[(vt["cid"] == cid) & (vt["sub"] == sub)]
    return float(m[f"v_{spec}"].iloc[0]) if len(m) else float("nan")


# =====================================================================================================================
# apply_upgrade, the planned schedule and the basis half-life (interface for Phase 3b and the final report)
# =====================================================================================================================
@lru_cache(maxsize=1)
def _token_index() -> dict:
    c = load_constraints()
    key_of = dict(zip(c["cid"], c["constraint_name"]))
    cids_of: dict = {}
    for cid, k in key_of.items():
        cids_of.setdefault(k, []).append(cid)
    typ = dict(zip(c["constraint_name"], c["element_type"]))
    ids = {r["upgrade_id"]: r["element_key"] for _, r in select_elements().iterrows()}
    return dict(key_of=key_of, cids_of=cids_of, typ=typ, ids=ids)


def residual_for(element_type: str, kind: str | None = None) -> float:
    """Shadow-price scale left after an upgrade: 0 = fully relieved. kind: 'wires' (line / transformer), 'gtc', 'full' (0.0); default by element type."""
    if kind == "full":
        return 0.0
    if kind == "gtc" or (kind is None and element_type == "gtc"):
        return RESIDUAL_MU_SCALE_GTC
    return RESIDUAL_MU_SCALE_WIRES


def apply_upgrade(cid_list, residual: float | None = None, expand_element: bool = True) -> dict:
    """Turn upgrades into per-constraint shadow-price scale factors: returns {cid: mu_scale}, where the shadow prices of that constraint are multiplied by
    mu_scale after the upgrade (0.0 = fully relieved).

    `cid_list` items may be a constraint id ('6437__F|DMTSCOS5'), a monitored element ('6437__F', the constraint_name) or an upgrade id ('U01' from
    data/upgrades.csv). With expand_element=True (default) every contingency constraint of the element is scaled, because the upgrade is physical
    (a reconductor relieves all contingencies); with False only the cids given are returned. `residual` overrides the default (RESIDUAL_MU_SCALE_WIRES for
    lines and transformers, RESIDUAL_MU_SCALE_GTC for interfaces); pass 0.0 for a full relief. Unknown items raise ValueError."""
    ix = _token_index()
    if isinstance(cid_list, str):
        cid_list = [cid_list]
    out: dict[str, float] = {}
    unknown = []
    for tok in cid_list:
        key = None
        if tok in ix["ids"]:
            key, members = ix["ids"][tok], ix["cids_of"][ix["ids"][tok]]
        elif tok in ix["key_of"]:
            key = ix["key_of"][tok]
            members = ix["cids_of"][key] if expand_element else [tok]
        elif tok in ix["cids_of"]:
            key, members = tok, ix["cids_of"][tok]
        else:
            unknown.append(tok)
            continue
        scale = float(residual) if residual is not None else residual_for(ix["typ"][key])
        for cid in members:
            out[cid] = scale if cid not in out else min(out[cid], scale)
    if unknown:
        raise ValueError(f"apply_upgrade: unknown constraint / element / upgrade ids: {unknown[:5]}")
    return out


@lru_cache(maxsize=1)
def planned_schedule_all() -> pd.DataFrame:
    c = load_constraints()
    d = c[c["upgrade_project_id"].notna()].copy()
    out = pd.DataFrame({
        "cid": d["cid"], "element_key": d["constraint_name"], "element_type": d["element_type"], "project_id": d["upgrade_project_id"],
        "project_name": d["upgrade_project_name"], "isd": d["upgrade_in_service_date"].map(parse_isd), "status": d["upgrade_status"],
        "conf": d["upgrade_match_conf"], "method": d["upgrade_match_method"], "source": "constraints.csv"})
    keys = set(PLANNED_REVIEW)
    out = out[~out["element_key"].isin(keys)]
    extra = []
    for key, rev in PLANNED_REVIEW.items():
        etype = c.loc[c["constraint_name"] == key, "element_type"].iloc[0]
        if rev["planned"] in ("yes", "likely"):
            for cid in c.loc[c["constraint_name"] == key, "cid"]:
                extra.append(dict(cid=cid, element_key=key, element_type=etype, project_id=rev["ids"], project_name="", isd=parse_isd(rev["isd"]),
                                  status=rev["status"], conf=rev["conf"], method="review", source="PLANNED_REVIEW"))
    out = pd.concat([out, pd.DataFrame(extra)], ignore_index=True)
    out["residual"] = [residual_for(t) for t in out["element_type"]]
    return out.reset_index(drop=True)


def planned_schedule(min_conf: float = UPGRADE_CONF_MIN) -> pd.DataFrame:
    """Planned-upgrade clock for every constraint with a project match at or above min_conf: cid, element, project id, in-service date (NaT if none),
    status, match confidence, default residual. A match is a name / bus overlap, not proof of relief (phase 2 notes); rows with status in_service are already
    in the market data. Projects without a date (RTP 'identified') carry NaT and never take effect in basis_half_life."""
    s = planned_schedule_all()
    return s[s["conf"] >= min_conf].reset_index(drop=True)


def basis_half_life(v_by_cid: dict, schedule: pd.DataFrame | None = None, t0: pd.Timestamp = AS_OF, horizon_years: float = 15.0,
                    residual: float | None = None, fraction: float = 0.5, include_in_service: bool = False,
                    undated_lead_months: float | None = None) -> dict:
    """Years from t0 until a site's congestion value V(t) = sum_k v_k x m_k(t) falls to `fraction` of V(t0), where m_k(t) is the product of the residuals of the planned
    upgrades of constraint k that are in service by t (dates before t0 count as t0). inf if it never does within the horizon. v_by_cid: {cid: $/kW-yr}.
    Projects that are matched but undated (RTP "identified") never take effect unless undated_lead_months is given: they then land that many months after t0."""
    sch = planned_schedule() if schedule is None else schedule
    if not include_in_service:
        sch = sch[sch["status"] != "in_service"]          # already in the market data behind v_by_cid
    sch = sch[sch["cid"].isin(v_by_cid)].copy()
    if undated_lead_months is not None:
        sch["isd"] = sch["isd"].fillna(t0 + pd.Timedelta(days=30.4375 * undated_lead_months))
    sch = sch[sch["isd"].notna()].copy()
    sch["t"] = sch["isd"].clip(lower=t0)
    sch = sch[sch["t"] <= t0 + pd.Timedelta(days=365.25 * horizon_years)].sort_values(["t", "cid"])
    v0 = float(sum(v_by_cid.values()))
    m = {cid: 1.0 for cid in v_by_cid}
    path, half = [(t0, v0)], float("inf")
    for t, g in sch.groupby("t"):
        for r in g.itertuples():
            m[r.cid] *= float(residual) if residual is not None else float(r.residual)
        v = float(sum(v_by_cid[cid] * m[cid] for cid in v_by_cid))
        path.append((t, v))
        if half == float("inf") and v <= fraction * v0:
            half = (t - t0).days / 365.25
    return dict(half_life_years=half, v0=v0, path=path)


@lru_cache(maxsize=1)
def connection_lookup() -> dict:
    """sub -> (connection class, lead time months, all-in interconnection $M) from data/connection_costs.csv (Phase 3a, read only); {} if absent."""
    path = paths.REPO_DATA / "connection_costs.csv"
    if not path.exists():
        return {}
    d = pd.read_csv(path, usecols=["sub", "connection_class", "lead_time_months", "all_in_mid_musd"])
    return {r.sub: (r.connection_class, float(r.lead_time_months), float(r.all_in_mid_musd)) for r in d.itertuples()}


# =====================================================================================================================
# Non-wires alternative sizing: data/nwa_sizing.csv
# =====================================================================================================================
SIDE_LABEL = {"import": "import_discharge", "export": "export_charge"}


def build_nwa() -> pd.DataFrame:
    """One row per element x side x relief target x coverage target. Battery MW = relief / |SF| over sites on the side that relieves (discharging on the
    import side, charging on the export side), MWh = MW x D(coverage), cost from the NREL line, compared with the primary and the cheapest eligible wires option."""
    rows = []
    for st in element_study():
        el, limit = st["el"], st["limit"]["limit_mw"]
        if not np.isfinite(limit):
            continue                                            # interface never priced in the window and no planning limit: not sized
        ep = st["episodes"]
        dur = {x: duration_at_coverage(ep, x) for x in COVERAGE_TARGETS}
        cid_used, sites = element_site_data(el)
        kc = _kc(el["kv_class"])
        prim = st["primary"]
        ah = st["annual"]["annual_hours"]
        for side in ("import", "export"):
            ss = None
            if sites is not None:
                ss = sites[sites["side"] == side]
                ss = ss if len(ss) else None
            fb = fallback_sf(kc) if (ss is None and el["element_type"] == "line" and kc in (69, 138, 345)) else (np.nan, np.nan, 0)
            for rf in RELIEF_TARGETS:
                relief = rf * limit
                if ss is not None:
                    sz = size_for_relief(relief, ss)
                    sf_conf = "measured" if sz["basis"] in ("rt_reg", "dam_reg") else "attenuated"
                else:
                    a = fb[0] if side == "import" else fb[1]
                    if not np.isfinite(a):
                        continue
                    mw = relief / a
                    sz = dict(mw_required=mw, n_sites=int(math.ceil(mw / site_cap_mw(el["kv"]))), shortfall_mw=0.0, mw_best_site=mw, best_site="(assumed)",
                              best_site_kv=np.nan, best_site_sf=(-a if side == "import" else a), best_site_zone="", basis="fallback_assumed")
                    sf_conf = "assumed"
                mw_inst = sz["mw_required"] / BATT_AVAILABILITY
                cheap = cheapest_eligible(st["options"], rf)
                for x in COVERAGE_TARGETS:
                    d = dur[x]
                    h_need = max(d, BATT_MIN_HOURS) if np.isfinite(d) else np.nan
                    h = min(h_need, BATT_MAX_HOURS) if np.isfinite(h_need) else np.nan
                    capped = bool(np.isfinite(h_need) and h_need > BATT_MAX_HOURS + 1e-9)
                    cov_ach = coverage_at(ep, h) if np.isfinite(h) else np.nan
                    cost = {k: battery_capex_musd(mw_inst, h, m) for k, m in BATT_COST_MULT.items()}
                    spec = "2h" if (np.isfinite(h) and h <= 2.5) else "4h"
                    cc = connection_lookup().get(sz["best_site"])
                    per_site_mva = BATT_GSU_MVA_PER_MW * mw_inst / max(sz["n_sites"], 1)
                    site_kv = sz["best_site_kv"] if np.isfinite(sz["best_site_kv"]) else el["kv"]          # assumed site: the element's own voltage
                    big = bool(np.isfinite(site_kv) and site_kv >= 138 and per_site_mva >= BATT_LPT_MIN_MVA)
                    site_lead = BATT_FIS_MONTHS[1] + (max(BATT_BUILD_MONTHS[1], LPT_LEAD_MONTHS[0] + XFMR_INSTALL_MONTHS) if big else BATT_BUILD_MONTHS[1])
                    refv = ref_value(cid_used, sz["best_site"], spec) if cid_used else float("nan")
                    rowd = dict(
                        upgrade_id=el["upgrade_id"], element=el["element"], element_type=el["element_type"], kv=el["kv"], side=SIDE_LABEL[side],
                        relief_pct=rf, relief_mw=relief, limit_mw=limit, limit_basis=st["limit"]["limit_basis"], coverage_pct=x,
                        episode_hours=d, battery_hours=h, duration_capped=capped, coverage_achieved=cov_ach,
                        episode_basis=st["episode_basis"], n_episodes=len(ep),
                        sf_basis=sz["basis"], sf_confidence=sf_conf, sf_cid=(cid_used or ""), best_site=sz["best_site"], best_site_kv=sz["best_site_kv"],
                        best_site_zone=sz["best_site_zone"], best_site_sf=sz["best_site_sf"], mw_best_site_uncapped=sz["mw_best_site"], n_sites=sz["n_sites"],
                        mw_required=sz["mw_required"], mw_installed=mw_inst, mwh_installed=mw_inst * h if np.isfinite(h) else np.nan,
                        battery_cost_low_musd=cost["low"], battery_cost_mid_musd=cost["mid"], battery_cost_high_musd=cost["high"],
                        battery_usd_per_kw_mid=battery_usd_per_kw(h, BATT_COST_MULT["mid"]) if np.isfinite(h) else np.nan,
                        battery_lead_months=site_lead, battery_lead_basis=("FIS median + large-power-transformer delivery and installation (step-up of 100 MVA or more at 138 kV or above)" if big
                                            else "FIS median + build (ASSUMPTION); step-up below large-power-transformer size"),
                        site_connection_class=(cc[0] if cc else ""))
                    if prim is not None and np.isfinite(prim["mid_musd"]):
                        rowd.update(wires_primary_option=prim["option"], wires_primary_low_musd=prim["low_musd"], wires_primary_mid_musd=prim["mid_musd"],
                                    wires_primary_high_musd=prim["high_musd"],
                                    ratio_vs_primary_mid=cost["mid"] / prim["mid_musd"], ratio_vs_primary_best=cost["low"] / prim["high_musd"],
                                    ratio_vs_primary_worst=cost["high"] / prim["low_musd"],
                                    ratio_annualised_vs_primary=battery_annual_musd(cost["mid"]) / wires_annual_musd(prim["mid_musd"]))
                    else:
                        rowd.update(wires_primary_option=(prim["option"] if prim else ""), wires_primary_low_musd=np.nan, wires_primary_mid_musd=np.nan,
                                    wires_primary_high_musd=np.nan, ratio_vs_primary_mid=np.nan, ratio_vs_primary_best=np.nan, ratio_vs_primary_worst=np.nan,
                                    ratio_annualised_vs_primary=np.nan)
                    if cheap is not None:
                        rowd.update(wires_cheapest_option=cheap["option"], wires_cheapest_mid_musd=cheap["mid_musd"],
                                    ratio_vs_cheapest_mid=cost["mid"] / cheap["mid_musd"])
                    else:
                        rowd.update(wires_cheapest_option="", wires_cheapest_mid_musd=np.nan, ratio_vs_cheapest_mid=np.nan)
                    rar = x * ah if np.isfinite(ah) else np.nan
                    rowd.update(annual_binding_hours=ah, annual_relief_hours=rar, cycles_per_year=(rar / h if np.isfinite(rar) and np.isfinite(h) else np.nan),
                                ref_zone_energy_usd_kw_yr=REF_ZONE_ENERGY_USD_KW_YR, ref_congestion_usd_kw_yr=refv,
                                ref_market_rev_musd_yr=(mw_inst * (REF_ZONE_ENERGY_USD_KW_YR + (refv if np.isfinite(refv) else 0.0)) / 1000.0))
                    rows.append(rowd)
    return pd.DataFrame(rows)


def nwa_best_side(nwa: pd.DataFrame) -> pd.DataFrame:
    """For each element x relief x coverage the side with the smaller installed MW."""
    x = nwa[nwa["mw_installed"].notna()]
    idx = x.groupby(["upgrade_id", "relief_pct", "coverage_pct"])["mw_installed"].idxmin()
    return x.loc[idx].reset_index(drop=True)


# =====================================================================================================================
# who decides (text)
# =====================================================================================================================
def decider_text(el: pd.Series, prim: dict | None, planned: dict) -> tuple[str, str]:
    tsp = f"TSP ({planned['planned_tsp']})" if planned.get("planned_tsp") else "the TSP (not identified from public data)"
    if el["element_type"] == "gtc":
        return ("ERCOT planning (RTP / STEP) proposes, the PUCT approves the plan (Project 55718) and the TSPs build and own the lines under CCNs; costs go into TCOS and are "
                "recovered from load by 4CP-based rates. Storage cannot be bought as an alternative; a stability-limited interface can be raised by grid-forming storage only "
                "through the FIS stability study and ERCOT's quarterly stability assessment (PG 5.3.2.4.3, 5.3.5)."), f"{CITE_PLAN}; {CITE_CCN}; {CITE_COST}"
    if prim is None:
        return f"{tsp[0].upper() + tsp[1:]} decides and builds", CITE_PLAN
    if prim["option"] in ("reconductor", "rebuild"):
        return (f"{tsp[0].upper() + tsp[1:]} decides, builds, owns and recovers the cost through TCOS. A {prim['option']} on the existing right of way needs no CCN. If it only changes ratings it is a "
                f"neutral Tier 4 project that ERCOT does not review or endorse (NP 3.11.4.1(4), 3.11.4.4); a 345 kV reconductor of more than one mile is Tier 3 and needs RPG acceptance, "
                f"and a rebuild of aged infrastructure over $50M starts as Tier 3. A stakeholder can ask the TSP or raise the constraint in the RPG; it cannot order it."), f"{CITE_NOCCN}; {CITE_PLAN}; {CITE_COST}"
    if prim["option"] == "autotransformer":
        return (f"{tsp[0].upper() + tsp[1:]} decides, builds, owns and recovers the cost through TCOS. A 345/138 kV autotransformer addition is Tier 3 (under $50M) with RPG acceptance; substations need "
                f"no CCN. The large power transformer lead time, not the permit, sets the schedule."), f"NP 3.11.4.3(1)(c)(ii); 16 TAC 25.101(c)(2); {CITE_PLAN}; {CITE_COST}"
    return f"{tsp[0].upper() + tsp[1:]} through the RPG and a PUCT CCN", f"{CITE_CCN}; {CITE_PLAN}; {CITE_COST}"


BATTERY_OWNER_TEXT = (
    "A battery owner can submit the upgrade to the RPG (any stakeholder, NP 3.11.4.1) and can pay for a TSP-elected, resource-funded project (NP 3.11.4.11), but cannot own, "
    "order or schedule network upgrades. Under connect-and-manage it pays its own interconnection facilities above a $14M (138 kV or less) / $20M (above 138 kV) allowance "
    "(16 TAC 25.195(f)(3)); the TSP pays network upgrades. No PUCT or ERCOT rule lets a TSP buy transmission-level storage in place of a line: PURA 35.153 (SB 415) covers "
    "distribution reliability only, bars TDU ownership and caps contracts at 100 MW (rule 16 TAC 25.58 proposed in Project 59523); ERCOT's planning compares wires alternatives and "
    "dispatch of existing resources (PG 3.1.2.1, 3.1.3, 4.1.1.2), and a RAS or CMP cannot be a permanent fix (PG 4.1.1.2(2)(a)).")


# =====================================================================================================================
# data/upgrades.csv and data/upgrade_options.csv
# =====================================================================================================================
def build_upgrades(nwa: pd.DataFrame | None = None) -> tuple[pd.DataFrame, pd.DataFrame]:
    nwa = build_nwa() if nwa is None else nwa
    best = nwa_best_side(nwa) if len(nwa) else nwa
    urows, orows = [], []
    for st in element_study():
        el, pl, ln, lim, prim, ah = st["el"], st["planned"], st["length"], st["limit"], st["primary"], st["annual"]
        uid = el["upgrade_id"]
        for o in st["options"]:
            orows.append({k: o[k] for k in ("upgrade_id", "option", "label", "low_musd", "mid_musd", "high_musd", "formula", "uplift", "ccn_needed", "process",
                                            "lead_low", "lead_mid", "lead_high", "lead_basis", "uses_unverified_cell")} | dict(
                element=el["element"], is_primary=bool(prim is not None and o["option"] == prim["option"])))
        b = best[best["upgrade_id"] == uid] if len(best) else best
        cen = b[(np.isclose(b["relief_pct"], CENTRAL_RELIEF)) & (np.isclose(b["coverage_pct"], CENTRAL_COVERAGE))] if len(b) else b
        nw = dict(nwa_central_side="", nwa_central_site="", nwa_central_site_sf=np.nan, nwa_central_mw=np.nan, nwa_central_mwh=np.nan, nwa_central_hours=np.nan,
                  nwa_mw_min=np.nan, nwa_mw_max=np.nan, nwa_mwh_min=np.nan, nwa_mwh_max=np.nan, nwa_cost_min_musd=np.nan, nwa_cost_central_musd=np.nan,
                  nwa_cost_max_musd=np.nan, nwa_vs_wires_ratio=np.nan, nwa_vs_wires_ratio_best=np.nan, nwa_vs_wires_ratio_worst=np.nan,
                  nwa_vs_cheapest_ratio=np.nan, nwa_cheapest_wires_option="", nwa_annualised_ratio=np.nan, nwa_sf_basis="", nwa_sf_confidence="",
                  nwa_central_coverage_achieved=np.nan, nwa_central_duration_capped=False, nwa_lead_time_months=np.nan, nwa_episode_basis=st["episode_basis"])
        if len(cen):
            c = cen.iloc[0]
            nw.update(nwa_central_side=c["side"], nwa_central_site=c["best_site"], nwa_central_site_sf=c["best_site_sf"], nwa_central_mw=c["mw_installed"],
                      nwa_central_mwh=c["mwh_installed"], nwa_central_hours=c["battery_hours"], nwa_central_coverage_achieved=c["coverage_achieved"],
                      nwa_central_duration_capped=bool(c["duration_capped"]), nwa_lead_time_months=c["battery_lead_months"], nwa_cost_central_musd=c["battery_cost_mid_musd"],
                      nwa_vs_wires_ratio=c["ratio_vs_primary_mid"], nwa_vs_wires_ratio_best=c["ratio_vs_primary_best"], nwa_vs_wires_ratio_worst=c["ratio_vs_primary_worst"],
                      nwa_vs_cheapest_ratio=c["ratio_vs_cheapest_mid"], nwa_cheapest_wires_option=c["wires_cheapest_option"],
                      nwa_annualised_ratio=c["ratio_annualised_vs_primary"], nwa_sf_basis=c["sf_basis"], nwa_sf_confidence=c["sf_confidence"])
        if len(b):
            nw.update(nwa_mw_min=b["mw_installed"].min(), nwa_mw_max=b["mw_installed"].max(), nwa_mwh_min=b["mwh_installed"].min(), nwa_mwh_max=b["mwh_installed"].max(),
                      nwa_cost_min_musd=b["battery_cost_low_musd"].min(), nwa_cost_max_musd=b["battery_cost_high_musd"].max())
        decider, cite = decider_text(el, prim, pl)
        urows.append(dict(
            upgrade_id=uid, element=el["element"], element_key=el["element_key"], element_type=el["element_type"], kv=el["kv"], n_contingencies=el["n_cids"],
            main_cid=el["main_cid"], priority_rank=el["best_rank"], selection=el["selection"], from_sub=el["from_sub"], to_sub=el["to_sub"],
            length_km=ln["length_km"], length_mi=ln["length_mi"], length_basis=ln["length_basis"], limit_mw=lim["limit_mw"], limit_basis=lim["limit_basis"],
            limit_swing_pct=st["swing"], imm_rent_2025_musd=el["imm_rent_2025_musd"], imm_market=el["imm_market"],
            annual_binding_hours=ah["annual_hours"], annual_hours_basis=ah["annual_hours_basis"],
            window_rt_hours=st["own"]["rt"].sum(), window_dam_hours=st["own"]["dam"].sum(),
            upgrade_type=(prim["option"] if prim else ""), upgrade_label=(prim["label"] if prim else ""), cost_formula=(prim["formula"] if prim else ""),
            cost_low_musd=(prim["low_musd"] if prim else np.nan), cost_mid_musd=(prim["mid_musd"] if prim else np.nan), cost_high_musd=(prim["high_musd"] if prim else np.nan),
            cost_uses_unverified_cell=(prim["uses_unverified_cell"] if prim else False), usd_year=USD_YEAR,
            imm_rent_to_annual_wires_cost=((el["imm_rent_2025_musd"] / wires_annual_musd(prim["mid_musd"]))
                                           if prim and np.isfinite(prim["mid_musd"]) and np.isfinite(el["imm_rent_2025_musd"]) else np.nan),
            lead_time_low_months=(prim["lead_low"] if prim else np.nan), lead_time_mid_months=(prim["lead_mid"] if prim else np.nan),
            lead_time_high_months=(prim["lead_high"] if prim else np.nan), lead_time_basis=(prim["lead_basis"] if prim else ""),
            other_options=" | ".join(f"{o['option']} {fm(o['mid_musd'])} (uplift {o['uplift']:.0%})" for o in st["options"] if prim is None or o["option"] != prim["option"]
                                    if np.isfinite(o["mid_musd"])),
            planned_yes_no=("yes" if pl["planned"] in ("yes", "likely") else "no"),
            planned=pl["planned"], planned_id=pl["planned_id"], planned_name=pl["planned_name"], planned_isd=pl["planned_isd"], planned_status=pl["planned_status"],
            planned_match_conf=pl["planned_conf"], planned_tsp=pl["planned_tsp"], planned_note=pl["planned_note"], planned_source=pl["planned_source"],
            **nw, decider=decider, battery_owner_role=BATTERY_OWNER_TEXT, citation=cite))
    return pd.DataFrame(urows), pd.DataFrame(orows)


# =====================================================================================================================
# First-order rerank: data/upgrade_rerank_first_order.csv
# =====================================================================================================================
def _not_rankable_reason(el: pd.Series) -> str:
    sf = load_shift_factors()
    m = sf[sf["cid"].isin(el["cids"])]
    if m.empty:
        return "no shift factor of any kind for this element's constraints (never priced in the window, or not fitted)"
    meths = set(m["method"])
    if meths == {"dcpf_proxy"}:
        return "only the DC power-flow proxy exists (never rank on proxy magnitudes, phase 2 gate B); the element never bound in the real-time week"
    return "only an attenuated (+shrunk) day-ahead fit exists: sign-only, magnitude not used"


def build_rerank() -> pd.DataFrame:
    """For each upgrade (and for PLANNED_ALL) and each spec: the sites that lose the most congestion value, and the sites that move up inside the post-upgrade
    top TOP_K. v0 = sum over constraints of V[n,k]; v1 = the same with the upgrade's constraints scaled by apply_upgrade; ranks are among sites with v0 >= V_MIN."""
    vt = value_table()
    lim = sph_table().set_index("cid")["limit_mw"]
    rows: list[dict] = []
    scenarios = []
    for st in element_study():
        el = st["el"]
        scenarios.append((el["upgrade_id"], el["element"], el["cids"], el))
    sch = planned_schedule()
    sch = sch[sch["isd"].notna() & (sch["status"] != "in_service")]
    scenarios.append(("PLANNED_ALL", "every constraint with a planned-project match >= 0.5, a date and not yet in service (planned_schedule())", sorted(set(sch["cid"])), None))
    for spec in CAPTURE_FACTOR:
        vcol = f"v_{spec}"
        g0 = vt.groupby("sub").agg(v0=(vcol, "sum"), zone=("zone8", "first"), kv_max=("kv_max", "max"), match_conf=("match_conf", "max"))
        universe = g0[g0["v0"] >= V_MIN_USD_KW_YR].copy()
        universe["rank0"] = universe["v0"].rank(ascending=False, method="first").astype(int)
        for uid, label, cids, el in scenarios:
            scales = apply_upgrade(cids, expand_element=False) if cids else {}
            hit = vt["cid"].isin(scales)
            if not hit.any():
                rows.append(dict(scenario=uid, element=label, spec=spec, group="not_rankable", order=0,
                                 note=_not_rankable_reason(el) if el is not None else "no constraint with a usable SF"))
                continue
            scale = vt["cid"].map(scales).fillna(1.0)
            vt1 = vt.assign(v1=vt[vcol] * scale, loss=vt[vcol] * (1 - scale))
            g1 = vt1.groupby("sub").agg(v1=("v1", "sum"), loss=("loss", "sum"))
            u = universe.join(g1, how="left").fillna({"v1": 0.0, "loss": 0.0})
            u["rank1"] = u["v1"].rank(ascending=False, method="first").astype(int)
            u["rank_change"] = u["rank0"] - u["rank1"]
            top = vt1.loc[hit].sort_values("loss", ascending=False).drop_duplicates("sub").set_index("sub")
            n_ranked = len(u)
            losers = u[u["loss"] > 0].sort_values(["loss", "v0"], ascending=False).head(N_LOSERS)
            risers = u[(u["rank1"] <= RISER_UNIVERSE) & (u["rank_change"] > 0)].sort_values(["rank_change", "v1"], ascending=False).head(N_RISERS)
            if risers.empty:
                rows.append(dict(scenario=uid, element=label, spec=spec, group="riser_none", order=0, n_ranked=n_ranked,
                                 note=f"no site in the post-upgrade top {RISER_UNIVERSE} moves up"))
            for grp, df in (("loser", losers), ("riser", risers)):
                for i, (sub, r) in enumerate(df.iterrows(), start=1):
                    t = top.loc[sub] if (grp == "loser" and sub in top.index) else None
                    move = ("leaves_top" if (r["rank0"] <= TOP_K < r["rank1"]) else "enters_top" if (r["rank1"] <= TOP_K < r["rank0"])
                            else "stays_in_top" if r["rank1"] <= TOP_K else "")
                    sat = (SATURATION_RELIEF * float(lim.get(t["cid"], np.nan)) / abs(t["dsf"])) if (t is not None and abs(t["dsf"]) > 0) else np.nan
                    rows.append(dict(
                        scenario=uid, element=label, spec=spec, group=grp, order=i, sub=sub, zone=r["zone"], kv_max=r["kv_max"], match_conf=r["match_conf"],
                        side=(t["side"] if t is not None else ""), main_cid=(t["cid"] if t is not None else ""), sf_main=(t["sf"] if t is not None else np.nan),
                        dsf_main=(t["dsf"] if t is not None else np.nan), sf_method=(t["method"] if t is not None else ""),
                        sph_usd_mw_yr=(t["sph_usd_mw_yr"] if t is not None else np.nan), sph_basis=(t["sph_basis"] if t is not None else ""), sat_mw=sat,
                        v0_usd_kw_yr=r["v0"], v1_usd_kw_yr=r["v1"], delta_usd_kw_yr=r["v1"] - r["v0"], share_lost=(r["loss"] / r["v0"] if r["v0"] > 0 else np.nan),
                        rank0=int(r["rank0"]), rank1=int(r["rank1"]), rank_change=int(r["rank_change"]), top_move=move, n_ranked=n_ranked, note=""))
    return pd.DataFrame(rows)


# =====================================================================================================================
# Label sidecars and writers
# =====================================================================================================================
def _spec(rows):
    return [tuple(r) for r in rows]


UPGRADES_LABELS = _spec([
    ("upgrade_id", "DERIVED", "gridsite.upgrades.select_elements", "U01.. in priority-rank order; the id apply_upgrade accepts"),
    ("element", "DERIVED", "data/constraints.csv", "from-to station and kV as ERCOT names them"),
    ("element_key", "REAL", "ERCOT NP6-86-CD / NP4-191-CD; IMM tables", "monitored element code (constraint_name); apply_upgrade accepts it"),
    ("element_type", "DERIVED", "data/constraints.csv", "line / xfmr / gtc (heuristic from station names and end voltages)"),
    ("kv", "REAL", "ERCOT NP6-86-CD", "larger end voltage, rounded"),
    ("n_contingencies", "DERIVED", "data/constraints.csv", "number of monitored-element | contingency constraints of this element"),
    ("main_cid", "DERIVED", "data/constraints.csv", "the contingency with the largest allocated rent, else the largest window rent"),
    ("priority_rank", "DERIVED", "data/constraints.csv", "best priority_rank over the element's contingencies"),
    ("selection", "ASSUMPTION", "gridsite.upgrades", "why the element is in the table: priority_top (top 15), imm_structural (2025 IMM top-10 rent), window_top (best local window constraints)"),
    ("from_sub", "DERIVED", "data/constraints.csv", "aliased NP4-160 substation; blank if the end has no placement"),
    ("to_sub", "DERIVED", "data/constraints.csv", ""),
    ("length_km", "DERIVED", "data/lines.csv; planned_projects.csv; substations.csv", "line route length; see length_basis. Blank for transformers and interfaces"),
    ("length_mi", "DERIVED", "length_km / 1.609344", "quantity used in the per-mile cost formulas"),
    ("length_basis", "DERIVED", "gridsite.upgrades.element_length", "lines.csv (OSM / HIFLD, same pair and voltage) > circuit-miles of a strongly matched planned project > station distance x 1.25 > class median (lines.csv)"),
    ("limit_mw", "REAL", "ERCOT NP6-86-CD / NP4-191-CD", "median SCED / DAM limit over binding rows; class median (DERIVED) where the element never priced in the window, see limit_basis"),
    ("limit_basis", "DERIVED", "gridsite.upgrades.element_limit", ""),
    ("limit_swing_pct", "DERIVED", "derived/shadow_sced.parquet", "(q95 - q05) / median of the SCED limit: a visible swing means ratings already follow ambient temperature (ERCOT 'Dynamic Ratings', NP section 2)"),
    ("imm_rent_2025_musd", "REAL", "Potomac Economics IMM Dec 2025 monthly report via data/constraints_annual.csv", "element-level 2025 congestion rent"),
    ("imm_market", "REAL", "data/constraints_annual.csv", "RT or DA; the two are never added"),
    ("annual_binding_hours", "DERIVED", "IMM 2025 hours; IMM rent x median hours per $M; or window x 365 / days x WINDOW_PERSISTENCE", "see annual_hours_basis"),
    ("annual_hours_basis", "DERIVED", "gridsite.upgrades.annual_hours", ""),
    ("window_rt_hours", "DERIVED", "derived/shadow_sced.parquet", "binding hours in the RT window (3-10 Oct and 25 Sep 2026), union over the contingencies, shadow price >= EPISODE_MIN_SHADOW"),
    ("window_dam_hours", "DERIVED", "derived/shadow_dam.parquet", "binding hours in the 31 DAM days, same rule"),
    ("upgrade_type", "ASSUMPTION", "gridsite.upgrades.primary_option", "the structural fix a TSP would most likely file: reconductor (lines; a rebuild where the matched planned project is a rebuild), parallel autotransformer (transformers), the 765 kV plan (WESTEX)"),
    ("upgrade_label", "DERIVED", "gridsite.upgrades.wires_options", ""),
    ("cost_formula", "DERIVED", "data/unit_costs.csv x quantity", "mid-case formula; low / high use the low / high unit-cost cells, widened to -20% / +50% where the cells are narrower"),
    ("cost_low_musd", "DERIVED", "MISO MTEP24 Texas rows in data/unit_costs.csv", "$M, 2024 dollars, not escalated"),
    ("cost_mid_musd", "DERIVED", "MISO MTEP24 Texas rows in data/unit_costs.csv", "$M, 2024 dollars"),
    ("cost_high_musd", "DERIVED", "MISO MTEP24 Texas rows in data/unit_costs.csv", "$M, 2024 dollars"),
    ("cost_uses_unverified_cell", "DERIVED", "data/unit_costs.csv note column", "True if a unit-cost cell used carries UNVERIFIED in its note (garbled MISO table layout)"),
    ("imm_rent_to_annual_wires_cost", "ASSUMPTION", "imm_rent_2025_musd / (cost_mid_musd x 11.5%)", "how many times one year of 2025 congestion rent covers one year of the wires fix; rent is not consumer savings, but PUCT's congestion test (16 TAC 25.101(b)(3)(A)(i)) compares a saving with the annual revenue requirement"),
    ("usd_year", "REAL", "MISO MTEP24; NREL 2025 update", "price year of every dollar figure"),
    ("lead_time_low_months", "DERIVED", "TPIT July 2026 base (planned projects), quartiles", "months from RPG submittal to projected in service; autotransformers floored by transformer lead time"),
    ("lead_time_mid_months", "DERIVED", "TPIT July 2026 base", "median"),
    ("lead_time_high_months", "DERIVED", "TPIT July 2026 base", "third quartile"),
    ("lead_time_basis", "DERIVED", "gridsite.upgrades.lead_months", ""),
    ("other_options", "DERIVED", "gridsite.upgrades.wires_options", "the other wires options with mid cost and typical rating uplift; all are in data/upgrade_options.csv"),
    ("planned_yes_no", "DERIVED", "planned column", "yes if planned is yes or likely (a project that names the element, match 0.5 or above, or reviewed); no otherwise, including identified (RTP, no date), adjacent and hint"),
    ("planned", "DERIVED", "data/constraints.csv upgrade_* and PLANNED_REVIEW", "yes (match >= 0.8 or reviewed) / likely (>= 0.5) / identified (RTP, no date) / adjacent or hint (nearby project, not the element) / done / no"),
    ("planned_id", "REAL", "data/planned_projects.csv", "TPIT / RTP / RPG / PBRP project id(s)"),
    ("planned_name", "REAL", "data/planned_projects.csv", ""),
    ("planned_isd", "REAL", "data/planned_projects.csv", "in-service date as TPIT / RPG print it (a year, a month or a day)"),
    ("planned_status", "REAL", "data/planned_projects.csv", ""),
    ("planned_match_conf", "ASSUMPTION", "gridsite.shiftfactors.MATCH_CONF", "a match is a name or bus overlap, not proof of relief"),
    ("planned_tsp", "REAL", "data/planned_projects.csv", "owner of the matched project"),
    ("planned_note", "ASSUMPTION", "analyst review of data/planned_projects.csv", "what the project is and why it is or is not the monitored element"),
    ("planned_source", "DERIVED", "gridsite.upgrades.planned_for", ""),
    ("nwa_central_side", "DERIVED", "data/nwa_sizing.csv", "the side with the smaller battery at the central scenario (10% relief, 80% coverage)"),
    ("nwa_central_site", "DERIVED", "data/shift_factors.csv", "best site on that side"),
    ("nwa_central_site_sf", "DERIVED", "derived/sf_long.parquet via data/shift_factors.csv", "its shift factor (sign convention PLAN.md: negative = import side)"),
    ("nwa_central_mw", "DERIVED", "relief MW / |SF|, per-site cap, / availability", "installed MW at the central scenario"),
    ("nwa_central_mwh", "DERIVED", "MW x D(coverage)", "installed MWh at the central scenario"),
    ("nwa_central_hours", "DERIVED", "episode lengths in the RT / DAM windows", "duration priced (1 to 10 h)"),
    ("nwa_central_coverage_achieved", "DERIVED", "coverage_at(episodes, nwa_central_hours)", "share of binding hours that duration actually covers (below 0.80 when the duration is capped at 10 h)"),
    ("nwa_central_duration_capped", "DERIVED", "episode_hours > 10 h", "True: binding is too persistent for storage to reach the coverage target"),
    ("nwa_lead_time_months", "DERIVED", "FIS median + build or transformer delivery", "battery lead time at the central scenario, months"),
    ("nwa_mw_min", "DERIVED", "data/nwa_sizing.csv", "smallest installed MW over the 9 relief x coverage scenarios (best side)"),
    ("nwa_mw_max", "DERIVED", "data/nwa_sizing.csv", "largest"),
    ("nwa_mwh_min", "DERIVED", "data/nwa_sizing.csv", ""),
    ("nwa_mwh_max", "DERIVED", "data/nwa_sizing.csv", ""),
    ("nwa_cost_min_musd", "DERIVED", "NREL/TP-6A40-93281, 2026 low case", "smallest scenario at the low cost case"),
    ("nwa_cost_central_musd", "DERIVED", "NREL/TP-6A40-93281, 2026 mid case", "central scenario"),
    ("nwa_cost_max_musd", "DERIVED", "NREL/TP-6A40-93281, 2026 high case", "largest scenario at the high cost case"),
    ("nwa_vs_wires_ratio", "DERIVED", "battery mid / primary wires mid", "above 1 means the battery costs more than the wires fix"),
    ("nwa_vs_wires_ratio_best", "DERIVED", "battery low / wires high", ""),
    ("nwa_vs_wires_ratio_worst", "DERIVED", "battery high / wires low", ""),
    ("nwa_vs_cheapest_ratio", "DERIVED", "battery mid / cheapest wires option whose typical uplift covers the relief", "the real competitor is often DLR / AAR or a reconductor"),
    ("nwa_cheapest_wires_option", "DERIVED", "gridsite.upgrades.cheapest_eligible", ""),
    ("nwa_annualised_ratio", "ASSUMPTION", "8% / 15 y + 4% FOM for the battery, 11.5% fixed charge for wires", "capex ratio re-expressed per year of life"),
    ("nwa_sf_basis", "DERIVED", "data/shift_factors.csv", "rt_reg / dam_reg / dam_reg_unshrunk / fallback_assumed"),
    ("nwa_sf_confidence", "DERIVED", "gridsite.upgrades.build_nwa", "measured (regression SF), attenuated (DAM SF divided by its shrink), assumed (class median, no regression SF)"),
    ("nwa_episode_basis", "DERIVED", "derived/shadow_sced.parquet, shadow_dam.parquet", "own_rt / own_dam episodes, or the pooled class of elements that did bind"),
    ("decider", "REAL", "NP 3.11.4; PG 3.1; 16 TAC 25.101, 25.192, 25.195", "who decides and who pays"),
    ("battery_owner_role", "REAL", "NP 3.11.4.1, 3.11.4.11; 16 TAC 25.195(f); PURA 35.153; PG 3.1.2.1, 4.1.1.2", "what a battery owner can and cannot do"),
    ("citation", "REAL", "see notes/phase4.md section 6", "rule references"),
])

OPTIONS_LABELS = _spec([
    ("upgrade_id", "DERIVED", "gridsite.upgrades.select_elements", ""),
    ("option", "ASSUMPTION", "gridsite.upgrades.wires_options", "reconductor / rebuild / new_parallel_line / dlr_aar / autotransformer / substation_rebuild / plan_765kv"),
    ("label", "DERIVED", "gridsite.upgrades.wires_options", ""),
    ("low_musd", "DERIVED", "MISO MTEP24 Texas rows in data/unit_costs.csv", "$M 2024; min(low cells, mid x 0.8)"),
    ("mid_musd", "DERIVED", "MISO MTEP24 Texas rows in data/unit_costs.csv", "$M 2024"),
    ("high_musd", "DERIVED", "MISO MTEP24 Texas rows in data/unit_costs.csv", "$M 2024; max(high cells, mid x 1.5)"),
    ("formula", "DERIVED", "data/unit_costs.csv x quantity", "mid-case arithmetic"),
    ("uplift", "ASSUMPTION", "DOE 2019 DLR report; MISO; judgement", "typical rating uplift; the option can serve a relief target only if uplift >= target"),
    ("ccn_needed", "REAL", "16 TAC 25.101(b)(3), (c)", "True for a new line"),
    ("process", "REAL", "NP 3.11.4.3; 16 TAC 25.101", "RPG tier and permit path"),
    ("lead_low", "DERIVED", "TPIT July 2026 base; DOE / NLR; ASSUMPTION for DLR", "months, first quartile"),
    ("lead_mid", "DERIVED", "TPIT July 2026 base; DOE / NLR; ASSUMPTION for DLR", "months, median"),
    ("lead_high", "DERIVED", "TPIT July 2026 base; DOE / NLR; ASSUMPTION for DLR", "months, third quartile"),
    ("lead_basis", "DERIVED", "gridsite.upgrades.lead_months", ""),
    ("uses_unverified_cell", "DERIVED", "data/unit_costs.csv note column", ""),
    ("element", "DERIVED", "data/constraints.csv", ""),
    ("is_primary", "ASSUMPTION", "gridsite.upgrades.primary_option", "the option quoted in data/upgrades.csv"),
])

NWA_LABELS = _spec([
    ("upgrade_id", "DERIVED", "gridsite.upgrades.select_elements", ""),
    ("element", "DERIVED", "data/constraints.csv", ""),
    ("element_type", "DERIVED", "data/constraints.csv", ""),
    ("kv", "REAL", "ERCOT NP6-86-CD", ""),
    ("side", "DERIVED", "PLAN.md sign convention", "import_discharge: sites with SF < 0, the battery discharges to cut the flow; export_charge: sites with SF > 0, it charges"),
    ("relief_pct", "ASSUMPTION", "brief: named scenarios", "relief target as a share of the median limit (market data cannot show the relief needed: SCED keeps flow at or below the limit)"),
    ("relief_mw", "DERIVED", "relief_pct x limit_mw", "flow relief the battery must deliver whenever the element binds"),
    ("limit_mw", "REAL", "ERCOT NP6-86-CD / NP4-191-CD", "median limit; see limit_basis"),
    ("limit_basis", "DERIVED", "gridsite.upgrades.element_limit", ""),
    ("coverage_pct", "ASSUMPTION", "brief", "share of binding hours that must be covered"),
    ("episode_hours", "DERIVED", "derived/shadow_sced.parquet, shadow_dam.parquet", "D(x): smallest duration such that relief sustained for at most D hours from the start of each binding episode covers x of all binding hours"),
    ("battery_hours", "ASSUMPTION", "min(max(episode_hours, BATT_MIN_HOURS), BATT_MAX_HOURS)", "duration of the battery (MWh / MW), between 1 and 10 h"),
    ("duration_capped", "DERIVED", "episode_hours > 10 h", "True: the binding is too persistent for storage to reach the coverage target; see coverage_achieved"),
    ("coverage_achieved", "DERIVED", "coverage_at(episodes, battery_hours)", "share of binding hours covered with the battery_hours actually priced"),
    ("episode_basis", "DERIVED", "gridsite.upgrades.episode_basis_for", "own_rt / own_dam: this element's episodes; pooled_*: the class of elements that did bind in the window"),
    ("n_episodes", "DERIVED", "gridsite.upgrades.episode_hours", ""),
    ("sf_basis", "DERIVED", "data/shift_factors.csv", "rt_reg / dam_reg / dam_reg_unshrunk (DAM SF divided by shrink) / fallback_assumed"),
    ("sf_confidence", "DERIVED", "gridsite.upgrades.build_nwa", "measured / attenuated / assumed"),
    ("sf_cid", "DERIVED", "data/constraints.csv", "the contingency constraint whose shift factors were used"),
    ("best_site", "DERIVED", "data/shift_factors.csv", "substation with the largest |SF| on the relieving side (69 kV and above, no weak / implausible rows)"),
    ("best_site_kv", "REAL", "data/substations.csv", ""),
    ("best_site_zone", "REAL", "data/substations.csv", ""),
    ("best_site_sf", "DERIVED", "derived/sf_long.parquet", "shift factor (negative = import side), after the shrink correction and before the cap"),
    ("mw_best_site_uncapped", "DERIVED", "relief_mw / |SF|", "MW if one site carried all of it, ignoring the per-site cap"),
    ("n_sites", "DERIVED", "gridsite.upgrades.size_for_relief", "sites needed with the per-site MW cap"),
    ("mw_required", "DERIVED", "greedy over sites by |SF|, capped per site", "MW that delivers the relief"),
    ("mw_installed", "ASSUMPTION", "mw_required / BATT_AVAILABILITY", "MW to install"),
    ("mwh_installed", "DERIVED", "mw_installed x battery_hours", "usable MWh"),
    ("battery_cost_low_musd", "DERIVED", "NREL/TP-6A40-93281 Fig 2 and Table 2 (2026 low)", "$M 2024: MW x (241 $/kWh x hours + 372 $/kW) x 0.77"),
    ("battery_cost_mid_musd", "DERIVED", "NREL/TP-6A40-93281 (2026 mid)", "x 0.92"),
    ("battery_cost_high_musd", "DERIVED", "NREL/TP-6A40-93281 (2026 high)", "x 1.10"),
    ("battery_usd_per_kw_mid", "DERIVED", "NREL/TP-6A40-93281", "$/kW at this duration"),
    ("battery_lead_months", "DERIVED", "FIS median 20.5 months + max(build 15, transformer 30 + 6) when the step-up (1.1 x MW per site) is 100 MVA or more at 138 kV or above, else + build 15", "months from the interconnection request to commercial operation; same logic as Phase 3a's connection_costs.csv"),
    ("battery_lead_basis", "DERIVED", "gridsite.upgrades.build_nwa", ""),
    ("site_connection_class", "DERIVED", "data/connection_costs.csv (Phase 3a)", "new_bay / new_switchyard / gen_tie at the best site; its cost is not added (NREL's cost already includes a permitting and interconnection line); blank where the best site is assumed"),
    ("wires_primary_option", "ASSUMPTION", "gridsite.upgrades.primary_option", ""),
    ("wires_primary_low_musd", "DERIVED", "data/upgrade_options.csv", ""),
    ("wires_primary_mid_musd", "DERIVED", "data/upgrade_options.csv", ""),
    ("wires_primary_high_musd", "DERIVED", "data/upgrade_options.csv", ""),
    ("ratio_vs_primary_mid", "DERIVED", "battery mid / primary wires mid", "above 1: the battery costs more than the wires fix"),
    ("ratio_vs_primary_best", "DERIVED", "battery low / wires high", ""),
    ("ratio_vs_primary_worst", "DERIVED", "battery high / wires low", ""),
    ("ratio_annualised_vs_primary", "ASSUMPTION", "(capex x (CRF 8%/15y + 4% FOM)) / (wires x 11.5%)", ""),
    ("wires_cheapest_option", "DERIVED", "gridsite.upgrades.cheapest_eligible", "cheapest option whose typical uplift covers relief_pct"),
    ("wires_cheapest_mid_musd", "DERIVED", "data/upgrade_options.csv", ""),
    ("ratio_vs_cheapest_mid", "DERIVED", "battery mid / cheapest eligible wires mid", ""),
    ("annual_binding_hours", "DERIVED", "IMM 2025 hours or window estimate", "see data/upgrades.csv annual_hours_basis"),
    ("annual_relief_hours", "DERIVED", "coverage_pct x annual_binding_hours", "hours a year the battery would be dispatched for relief"),
    ("cycles_per_year", "DERIVED", "annual_relief_hours / battery_hours", "rough full cycles a year; above about 700 the 2 h wear assumption breaks"),
    ("ref_zone_energy_usd_kw_yr", "REAL", "docs/research/arena/README.md", "ERCOT 5-year Adaptive learner, 2 h grid battery; reference only, never netted against cost"),
    ("ref_congestion_usd_kw_yr", "DERIVED", "data/upgrade_rerank_first_order.csv method", "first-order congestion value at the best site for this constraint; reference only"),
    ("ref_market_rev_musd_yr", "DERIVED", "mw_installed x (zone energy + congestion)", "$M a year the battery would also earn in the market: reference only, never netted"),
])

RERANK_LABELS = _spec([
    ("scenario", "DERIVED", "gridsite.upgrades.build_rerank", "upgrade id (U01..) or PLANNED_ALL"),
    ("element", "DERIVED", "data/constraints.csv", ""),
    ("spec", "REAL", "arena SPECS / gridsite.specs", "2h = arena 'grid', 4h = 'grid4h'"),
    ("group", "DERIVED", "gridsite.upgrades.build_rerank", "loser: biggest $/kW-yr loss; riser: biggest rank gain inside the post-upgrade top 200 (its own value is unchanged unless it shares a constraint); riser_none: no such site; not_rankable: no usable regression SF"),
    ("order", "DERIVED", "", "position inside the group"),
    ("sub", "REAL", "data/shift_factors.csv", "NP4-160 substation"),
    ("zone", "REAL", "data/settlement_points.csv", "arena load zone (NOIE zone where there is one)"),
    ("kv_max", "REAL", "data/substations.csv", ""),
    ("match_conf", "ASSUMPTION", "data/substations.csv", "placement confidence; the value does not depend on placement, the siting does"),
    ("side", "DERIVED", "PLAN.md sign convention", "side of the constraint that contributes most of the loss: import (price rises when it binds) or export"),
    ("main_cid", "DERIVED", "", "constraint that contributes most of the loss"),
    ("sf_main", "DERIVED", "derived/sf_long.parquet", ""),
    ("dsf_main", "DERIVED", "sf_main - zone mean", "the zone mean is the unweighted mean over the zone's buses of 60 kV and above"),
    ("sf_method", "DERIVED", "derived/sf_long.parquet", "rt_reg or non-attenuated dam_reg only"),
    ("sph_usd_mw_yr", "DERIVED", "IMM 2025 rent / limit, or window rent annualised", "annual shadow-price-hours of the main constraint, $/MW-yr (sum over a year of the shadow price in $/MWh x 1 h)"),
    ("sph_basis", "ASSUMPTION", "gridsite.upgrades.sph_table", "imm_rent_alloc (REAL rent) / window_rt / window_dam (x 365 / days x WINDOW_PERSISTENCE, capped at the 10th-ranked IMM rent / limit)"),
    ("sat_mw", "ASSUMPTION", "0.20 x limit / |dsf_main|", "battery MW that already relieves 20% of the main constraint: the per-kW values hold for roughly the first sat_mw only (price-taker limit, PROMPT item 25)"),
    ("v0_usd_kw_yr", "ASSUMPTION", "sum_k min(|dSF|, 1) x shadow-price-hours / 1000 x CAPTURE_FACTOR", "first-order congestion value before the upgrade; capture factor = 0.41 x cover of 2 h / 4 h"),
    ("v1_usd_kw_yr", "ASSUMPTION", "same, with the upgraded constraints scaled by apply_upgrade", "value after"),
    ("delta_usd_kw_yr", "ASSUMPTION", "v1 - v0", "negative = loses value"),
    ("share_lost", "ASSUMPTION", "(v0 - v1) / v0", ""),
    ("rank0", "ASSUMPTION", "rank of v0 among sites with v0 >= 0.5 $/kW-yr", "1 = best"),
    ("rank1", "ASSUMPTION", "rank of v1 in the same set", ""),
    ("rank_change", "ASSUMPTION", "rank0 - rank1", "positive = moves up"),
    ("top_move", "ASSUMPTION", "rank0 / rank1 against TOP_K = 50", "enters_top / leaves_top / stays_in_top"),
    ("n_ranked", "DERIVED", "", "sites ranked"),
    ("note", "DERIVED", "", "reason when not rankable"),
])


def _round_df(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        if out[c].dtype.kind == "f":
            nd = 4 if ("sf" in c or "ratio" in c or c.startswith("dsf") or c in ("uplift", "planned_match_conf", "match_conf", "limit_swing_pct", "share_lost")) else 3
            out[c] = out[c].round(nd)
    return out


INT_COLUMNS = ("order", "rank0", "rank1", "rank_change", "n_ranked", "n_sites", "n_episodes", "n_contingencies", "priority_rank", "usd_year")


def write_table(df: pd.DataFrame, path, labels: list) -> None:
    out = _round_df(df)
    for c in INT_COLUMNS:
        if c in out.columns:
            out[c] = out[c].round().astype("Int64")
    out.to_csv(path, index=False, encoding="utf-8")
    write_labels(path, labels, out)


# =====================================================================================================================
# Commands
# =====================================================================================================================
def run_upgrades(*args) -> pd.DataFrame:
    if "check" in args:
        return run_check()
    nwa = build_nwa()
    up, opt = build_upgrades(nwa)
    write_table(up, OUT_UPGRADES, UPGRADES_LABELS)
    write_table(opt, OUT_OPTIONS, OPTIONS_LABELS)
    return up


def run_nwa(*args) -> pd.DataFrame:
    nwa = build_nwa()
    write_table(nwa, OUT_NWA, NWA_LABELS)
    return nwa


def run_rerank(*args) -> pd.DataFrame:
    rr = build_rerank()
    write_table(rr, OUT_RERANK, RERANK_LABELS)
    return rr


# =====================================================================================================================
# Self-check: recompute the derived constants and compare with the named ones
# =====================================================================================================================
def derive_cover() -> dict:
    """Value-weighted share of a constraint's shadow-price-hours reachable by the best d hours of each day (d = 1, 2, 4, 6), RT and DAM windows."""
    sced, dam = window_tables()
    out = {}
    for name, df, dt, dcol in (("rt", sced, SCED_DT_MIN / 60.0, "i15"), ("dam", dam, 1.0, "h")):
        b = df[df["shadow_price"] > 0].copy()
        b["hr"] = b[dcol] // 4 if dcol == "i15" else b[dcol]
        b["mu_dt"] = b["shadow_price"] * dt
        g = b.groupby(["cid", "date_cpt", "hr"])["mu_dt"].sum().reset_index()
        tot = g.groupby("cid")["mu_dt"].sum()
        keep = tot[tot > COVER_MIN_SPH].index
        g = g[g["cid"].isin(keep)].sort_values(["cid", "date_cpt", "mu_dt"], ascending=[True, True, False])
        g["rk"] = g.groupby(["cid", "date_cpt"]).cumcount()
        res = {}
        for d in (1, 2, 4, 6):
            top = g[g["rk"] < d].groupby("cid")["mu_dt"].sum()
            res[d] = float(top.sum() / tot[keep].sum())
        out[name] = res
    return out


def run_check(*args) -> None:
    print("== derived constants vs the named ones")
    cv = derive_cover()
    for d, name in ((2, "COVER_2H"), (4, "COVER_4H")):
        rt, dm = cv["rt"][d], cv["dam"][d]
        print(f"  {name}: window RT {rt:.3f}, DAM {dm:.3f}, mean {(rt + dm) / 2:.3f}; named {globals()[name]:.3f}")
    print(f"  CAPTURE_FACTOR = {ARENA_LEARNER_SHARE} x cover: 2h {CAPTURE_FACTOR['2h']:.3f}, 4h {CAPTURE_FACTOR['4h']:.3f}")
    E = all_elements().set_index("element_key")
    sk = E.loc["6437__F"]
    win_rt = float(sk["window_rent_rt_musd"]) * 365.0 / rt_window_days()
    print(f"  WINDOW_PERSISTENCE: Scurry-Knapp IMM RT rent {sk['imm_rent_2025_musd']:.1f} M / window RT rent annualised {win_rt:.1f} M = {sk['imm_rent_2025_musd'] / win_rt:.2f}; named {WINDOW_PERSISTENCE}")
    k, n = imm_hours_per_musd()
    print(f"  IMM binding hours per $M rent: median {k:.2f} over {n} elements")
    for kc in (69, 138, 345):
        a, b, n = fallback_sf(kc)
        print(f"  fallback best-site |SF| for {kc} kV lines: import {a:.3f}, export {b:.3f} (n={n})")
    print("  battery $/kW (2024$, mid case x0.92): " + ", ".join(f"{h:g} h {battery_usd_per_kw(h, BATT_COST_MULT['mid']):,.0f}" for h in (1, 2, 4, 6)))
    print(f"  4 h at the NREL base case (x1.00): {battery_usd_per_kw(4, 1.0):,.0f} $/kW = {battery_usd_per_kw(4, 1.0) / 4:,.0f} $/kWh (report: 334)")
    t = tpit_lead_table()
    print("  TPIT lead times (months, RPG submittal to in service), pipeline all-kV:")
    print(t[(t['basis'] == 'pipeline') & (t['kvc'] == 'all')].round(1).to_string(index=False))
    print("  TPIT completed projects (actual), all-kV:")
    print(t[(t['basis'] == 'completed') & (t['kvc'] == 'all')].round(1).to_string(index=False))
    print(f"  episode check: coverage_at(duration_at_coverage(x)) round trip on Scurry-Knapp RT episodes -> "
          + ", ".join(f"{x:.2f}: D={duration_at_coverage(all_episodes()['6437__F']['rt'], x):.2f} h cov={coverage_at(all_episodes()['6437__F']['rt'], duration_at_coverage(all_episodes()['6437__F']['rt'], x)):.3f}" for x in COVERAGE_TARGETS))
